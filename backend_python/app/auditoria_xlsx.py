"""Excel de auditoria: los datos de entrada de cada bodega x referencia como VALORES y todo el calculo del modelo (la hoja
"Modelo" de Descripcion y Formulas DDMRP.xlsx) hasta la cantidad sugerida como FORMULAS vivas de Excel, mas lo que calculo el
motor y una columna que dice si coinciden. Los parametros y los tramos de factores viven en la hoja Parametros: se pueden
cambiar ahi y todas las filas se recalculan (para probar "que pasaria si")."""
import io
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

AZUL = Font(color="1F4E9E")          # entradas
NEGRO = Font(color="000000")         # formulas
GRIS = Font(color="5F6368")          # lo que dijo el motor
NEG_B = Font(bold=True, color="FFFFFF")
F_ENT, F_FOR, F_MOT, F_REV = (PatternFill("solid", fgColor=c) for c in ("EAF1FB", "FFFFFF", "F1F3F4", "FFF4CE"))
H_ENT, H_FOR, H_MOT = (PatternFill("solid", fgColor=c) for c in ("1F4E9E", "4F7A28", "5F6368"))

# Parametros: fila en la hoja Parametros (columna B)
PARAMS = [
    ("mult_cota_max_promo", "Multiplicador de la cota máxima en promoción", 1.2),
    ("mult_cota_min_promo", "Multiplicador de la cota mínima en promoción", 1.1),
    ("adu_bajo_min", "ADU bajo: desde (ZRB = 0 si el ADU está entre desde y hasta)", 1),
    ("adu_bajo_max", "ADU bajo: hasta", 4),
    ("regla_disparo", "Se pide cuando: 1 flujo neto ≤ TOY aj · 2 ≤ TOR · 3 < TOG aj", 1),
    ("regla_objetivo", "Se repone hasta (regla general): 1 TOG aj · 2 TOY aj · 3 TOY aj + mitad de la zona verde", 1),
    ("regla_redondeo", "Múltiplo de despacho: 1 arriba · 2 cercano · 3 abajo", 1),
    ("sumar_vectorizacion", "Sumar la vectorización (1 sí, 0 no)", 1),
    ("incluir_transito", "Incluir el tránsito en el flujo neto (1 sí, 0 no)", 1),
    ("restar_demanda_calificada", "Restar la demanda calificada (1 sí, 0 no)", 1),
    ("cantidad_minima", "Cantidad mínima: lo sugerido menor a esto se descarta", 0),
    ("meses_historia", "Meses de venta de la ventana (informativo: define las columnas de ventas)", 7),
    ("meses_adu", "Últimos meses que entran al ADU (informativo: define el rango de la fórmula)", 4),
]
PR = {k: f"Parametros!$B${i + 2}" for i, (k, _, _) in enumerate(PARAMS)}
TRAMO0 = len(PARAMS) + 5   # primera fila de tramos en Parametros (1 encabezado + params + vacia + titulo + encabezado)


def _tramos_formula(tramos: list[dict], x: str, offset: int) -> str:
    """Replica los SI(Y(x>=desde; x<=hasta)...) del Excel: el primer tramo incluye 'desde', los demas son (desde, hasta]."""
    f = "0"
    for i in reversed(range(len(tramos))):
        r = offset + i
        cond = f"{x}{'>=' if i == 0 else '>'}Parametros!$C${r}"
        if tramos[i]["hasta"] is not None:
            cond += f",{x}<=Parametros!$D${r}"
        f = f"IF(AND({cond}),Parametros!$E${r},{f})"
    return f


def construir(filas: list[dict], meses: list[str], params: dict, factores: list[dict], ejecucion: dict, truncado: int = 0) -> bytes:
    n = len(meses)
    k = max(1, min(int(params.get("meses_adu", 4)), n))
    lt_t = sorted([f for f in factores if f["tipo"] == "LT"], key=lambda f: f["desde"])
    va_t = sorted([f for f in factores if f["tipo"] == "VAR"], key=lambda f: f["desde"])
    off_lt = TRAMO0
    off_va = off_lt + len(lt_t)

    # ---- definicion de columnas: (clave, titulo, descripcion, tipo, ancho, formato) ----
    cols: list[tuple] = []
    def col(key, titulo, desc, tipo, ancho=11, fmt=None): cols.append((key, titulo, desc, tipo, ancho, fmt))
    col("emp", "Empresa", "05 Funat · 21 Stores", "e", 8)
    col("bod", "Bodega", "Tienda destino", "e", 9)
    col("ref", "Referencia", "Código del producto", "e", 11)
    col("desc", "Descripción", "Nombre del producto", "e", 34)
    col("lab", "Laboratorio", "Proveedor / laboratorio", "e", 22)
    for i, m in enumerate(meses):
        col(f"v{i}", m[:4] + "-" + m[4:] + (" (mes en curso)" if i == n - 1 else ""), "Unidades vendidas en el mes (ventas por mes cargadas)", "e", 10, "#,##0")
    col("pron", "Pronóstico", "Valor estático que cargan los analistas (unidades/mes)", "e", 11, "#,##0.##")
    col("lt", "Lead time (días)", "Días entre el pedido y la llegada a la tienda (bitácora)", "e", 9)
    col("frec", "Frecuencia (días)", "Cada cuántos días se despacha a la tienda (bitácora)", "e", 9)
    col("ud", "Unidad de despacho", "Múltiplo en que se puede enviar", "e", 9)
    col("cmax", "Cota máxima", "Control de referencias", "e", 9, "0.0000")
    col("cmin", "Cota mínima", "Control de referencias", "e", 9, "0.0000")
    col("promo", "Promoción", "0 no · 1 sí · 2 variante GN (no aplica a TC ni DM)", "e", 9)
    col("vect", "Vectorización", "Unidades de derrame que se suman al pedido", "e", 10, "#,##0.##")
    col("inv", "Inventario", "Disponible en la tienda al correr el motor", "e", 10, "#,##0.##")
    col("tr", "En tránsito", "Traslados abiertos hacia la tienda", "e", 10, "#,##0.##")
    col("dc", "Demanda calificada", "Pedidos ya comprometidos (se resta del flujo neto)", "e", 10, "#,##0.##")
    col("cond", "Condenado (NxP)", "1 = no se despacha nada", "e", 10)
    col("robj", "Regla hasta dónde (item/lab.)", "0 = usa la regla general; 1, 2 o 3 = la regla propia del item o de su laboratorio al correr el motor", "e", 11)
    # --- bloque pronostico (BI..CB del Excel)
    col("bi", "Promedio ventas", "Promedio de los meses de la ventana (BI)", "f", 10, "0.00")
    col("bl", "Desv. est.", "Desviación estándar muestral de la ventana (BL)", "f", 10, "0.00")
    col("bm", "IV ventana", "Desviación / promedio (BM)", "f", 9, "0.000")
    col("bk", "ADU referencia", "Promedio + promedio × IV, redondeado arriba (BK)", "f", 10)
    col("ca", "Días inv. pronóstico", "Pronóstico / ADU referencia × 30 (CA)", "f", 10, "0.0")
    col("q", "Forecast ajustado", "Si <30 días: pronóstico/días×30 · si >60: pronóstico/días×60 · si no, el pronóstico (Q = CB)", "f", 10, "0.00")
    # --- bloque principal (R..AM)
    col("u", "Promedio ADU", f"Promedio de los últimos {k} meses y el forecast (U)", "f", 10, "0.00")
    col("x", "Desv. est. ADU", f"Desviación estándar de los últimos {k} meses y el forecast (X)", "f", 10, "0.00")
    col("y", "IV", "Desviación / promedio ADU (Y)", "f", 9, "0.000")
    col("w", "ADU", "Promedio ADU + promedio ADU × IV, redondeado arriba (W)", "f", 9)
    col("za", "ZA", "Zona amarilla = lead time × ADU / 30, redondeado arriba (AG)", "f", 8)
    col("flt", "Factor LT", "Factor según el tramo del lead time (Z) — tabla en Parametros", "f", 9, "0.00")
    col("zrb", "ZRB", "Zona roja base = ZA × factor LT (0 si el ADU es bajo) (AE)", "f", 8)
    col("fvar", "Factor variab.", "Factor según el tramo del IV (AB) — tabla en Parametros", "f", 9, "0.00")
    col("zrs", "ZRS", "Zona roja de seguridad = ZRB × factor de variabilidad (AF)", "f", 8)
    col("tor", "TOR", "Tope rojo = ZRB + ZRS (AI)", "f", 8)
    col("toy", "TOY", "Tope amarillo = TOR + ZA (AJ)", "f", 8)
    col("zv", "ZV", "Zona verde = máx(frecuencia × ADU/30 arriba, unidad de despacho, ZRB) (AH)", "f", 8)
    col("tog", "TOG", "Tope verde = TOY + ZV (AK)", "f", 8)
    col("pe", "Promo efectiva", "La promoción se anula en bodegas TC (tiendas de confianza) y DM (domicilios)", "f", 9)
    col("toyaj", "TOY ajustado", "Tope amarillo con cotas / promoción (AL)", "f", 10)
    col("togaj", "TOG ajustado", "Tope verde con cotas / promoción (AM)", "f", 10)
    # --- decision
    col("nfp", "Flujo neto", "Inventario + tránsito − demanda calificada (según Parametros)", "f", 10, "#,##0.##")
    col("zona", "Zona", "ROJA ≤ TOR · AMARILLA ≤ TOY aj · VERDE ≤ TOG aj · si no, EXCESO", "f", 10)
    col("pide", "¿Pide?", "1 si cumple la regla de disparo", "f", 7)
    col("falta", "Falta", "Objetivo − flujo neto. Objetivo según la regla: 1 TOG aj · 2 TOY aj · 3 TOY aj + mitad de la zona verde (arriba)", "f", 9, "#,##0.##")
    col("ude", "Unidad efectiva", "Unidad de despacho (mínimo 1)", "f", 8)
    col("base", "Pedido al múltiplo", "Falta redondeada al múltiplo de despacho", "f", 10)
    col("basem", "Pedido tras mínimo", "Se descarta si es menor a la cantidad mínima", "f", 10)
    col("vef", "Vectorización sumada", "La vectorización se suma tal cual", "f", 10)
    col("nec", "NECESIDAD (fórmula)", "Lo que pide el buffer = pedido + vectorización (0 si es NxP)", "f", 12, "#,##0.##")
    # --- motor
    col("m_adu", "ADU motor", "Lo que calculó el motor", "m", 9)
    col("m_toyaj", "TOY aj motor", "Lo que calculó el motor", "m", 10)
    col("m_togaj", "TOG aj motor", "Lo que calculó el motor", "m", 10)
    col("m_zona", "Zona motor", "Lo que calculó el motor", "m", 10)
    col("m_nec", "Necesidad motor", "Antes de repartir el CEDI", "m", 11, "#,##0.##")
    col("m_cedi", "Stock CEDI repartible", "Inventario de PT001 disponible para esta referencia al correr el motor (después de la reserva)", "m", 11, "#,##0.##")
    col("m_asig", "ASIGNADO motor", "Lo que quedó después de repartir el CEDI (lo que ve la matriz de despacho)", "m", 12, "#,##0.##")
    col("m_est", "Estado", "PENDIENTE · APROBADA · ENVIADA", "m", 11)
    col("m_apr", "Cant. aprobada", "Si se editó o aprobó otra cantidad", "m", 10, "#,##0.##")
    col("rev", "REVISIÓN", "OK si la fórmula reproduce ADU, topes, zona y necesidad del motor", "f", 11)

    idx = {c[0]: i + 1 for i, c in enumerate(cols)}
    C = lambda key, r: f"{L(idx[key])}{r}"
    R0 = 4                                   # primera fila de datos
    hist = lambda r: f"{C('v0', r)}:{C(f'v{n - 1}', r)}"
    ult = lambda r: f"{C(f'v{n - k}', r)}:{C(f'v{n - 1}', r)}"
    p = PR
    F = {
        "bi": lambda r: f"=AVERAGE({hist(r)})",
        "bl": lambda r: f"=STDEV({hist(r)})",
        "bm": lambda r: f"=IFERROR({C('bl', r)}/{C('bi', r)},0)",
        "bk": lambda r: f"=ROUNDUP({C('bi', r)}+{C('bi', r)}*{C('bm', r)},0)",
        "ca": lambda r: f"=IFERROR({C('pron', r)}/{C('bk', r)}*30,0)",
        "q": lambda r: f"=IF({C('ca', r)}=0,0,IF({C('ca', r)}<30,{C('pron', r)}/{C('ca', r)}*30,IF({C('ca', r)}>60,{C('pron', r)}/{C('ca', r)}*60,{C('pron', r)})))",
        "u": lambda r: f"=AVERAGE({ult(r)},{C('q', r)})",
        "x": lambda r: f"=STDEV({ult(r)},{C('q', r)})",
        "y": lambda r: f"=IFERROR({C('x', r)}/{C('u', r)},0)",
        "w": lambda r: f"=ROUNDUP({C('u', r)}+{C('u', r)}*{C('y', r)},0)",
        "za": lambda r: f"=ROUNDUP({C('lt', r)}*{C('w', r)}/30,0)",
        "flt": lambda r: "=" + _tramos_formula(lt_t, C('lt', r), off_lt),
        "zrb": lambda r: f"=IF(AND({C('w', r)}>={p['adu_bajo_min']},{C('w', r)}<={p['adu_bajo_max']}),0,ROUND({C('za', r)}*{C('flt', r)},0))",
        "fvar": lambda r: "=" + _tramos_formula(va_t, C('y', r), off_va),
        "zrs": lambda r: f"=ROUND({C('zrb', r)}*{C('fvar', r)},0)",
        "tor": lambda r: f"={C('zrb', r)}+{C('zrs', r)}",
        "toy": lambda r: f"={C('tor', r)}+{C('za', r)}",
        "zv": lambda r: f"=IF(OR({C('w', r)}=0,{C('toy', r)}=0),0,MAX(ROUNDUP({C('frec', r)}*{C('w', r)}/30,0),{C('ud', r)},{C('zrb', r)}))",
        "tog": lambda r: f"={C('toy', r)}+{C('zv', r)}",
        "pe": lambda r: f"=IF(AND({C('promo', r)}<>0,OR(LEFT({C('bod', r)},2)=\"TC\",LEFT({C('bod', r)},2)=\"DM\")),0,{C('promo', r)})",
        "toyaj": lambda r: (f"=IF(AND(AND({C('w', r)}>={p['adu_bajo_min']},{C('w', r)}<={p['adu_bajo_max']}),{C('pe', r)}=0),{C('toy', r)},"
                            f"IF({C('pe', r)}=1,ROUNDUP({C('tog', r)}*{C('cmax', r)}*{p['mult_cota_max_promo']}+{C('tor', r)}+{C('za', r)},0),"
                            f"ROUNDUP({C('tog', r)}*{C('cmax', r)}+{C('tor', r)}+{C('za', r)},0)))"),
        "togaj": lambda r: (f"=IF(AND(AND({C('w', r)}>={p['adu_bajo_min']},{C('w', r)}<={p['adu_bajo_max']}),{C('pe', r)}=0),{C('tog', r)},"
                            f"IF(OR({C('pe', r)}=1,AND({C('pe', r)}=2,LEFT({C('bod', r)},2)=\"GN\")),"
                            f"ROUNDUP({C('toyaj', r)}+{C('zv', r)}+{C('tog', r)}*{C('cmin', r)}*{p['mult_cota_min_promo']},0),"
                            f"ROUNDUP({C('toyaj', r)}+{C('zv', r)}+{C('tog', r)}*{C('cmin', r)},0)))"),
        "nfp": lambda r: f"={C('inv', r)}+IF({p['incluir_transito']}=1,{C('tr', r)},0)-IF({p['restar_demanda_calificada']}=1,{C('dc', r)},0)",
        "zona": lambda r: (f"=IF({C('togaj', r)}<=0,IF({C('nfp', r)}>0,\"EXCESO\",\"VERDE\"),IF({C('nfp', r)}<={C('tor', r)},\"ROJA\","
                           f"IF({C('nfp', r)}<={C('toyaj', r)},\"AMARILLA\",IF({C('nfp', r)}<={C('togaj', r)},\"VERDE\",\"EXCESO\"))))"),
        "pide": lambda r: (f"=IF(AND({C('togaj', r)}>0,IF({p['regla_disparo']}=2,{C('nfp', r)}<={C('tor', r)},"
                           f"IF({p['regla_disparo']}=3,{C('nfp', r)}<{C('togaj', r)},{C('nfp', r)}<={C('toyaj', r)}))),1,0)"),
        "falta": lambda r: (f"=IF({C('pide', r)}=1,CHOOSE(IF({C('robj', r)}>0,{C('robj', r)},{p['regla_objetivo']}),{C('togaj', r)},{C('toyaj', r)},"
                            f"ROUNDUP({C('toyaj', r)}+MAX({C('togaj', r)}-{C('toyaj', r)},0)/2,0))-{C('nfp', r)},0)"),
        "ude": lambda r: f"=MAX({C('ud', r)},1)",
        "base": lambda r: (f"=IF({C('falta', r)}>0,{C('ude', r)}*IF({p['regla_redondeo']}=2,ROUND({C('falta', r)}/{C('ude', r)},0),"
                           f"IF({p['regla_redondeo']}=3,ROUNDDOWN({C('falta', r)}/{C('ude', r)},0),ROUNDUP({C('falta', r)}/{C('ude', r)},0))),0)"),
        "basem": lambda r: f"=IF(AND({C('base', r)}>0,{C('base', r)}<{p['cantidad_minima']}),0,{C('base', r)})",
        "vef": lambda r: f"=IF({p['sumar_vectorizacion']}=1,{C('vect', r)},0)",
        "nec": lambda r: f"=IF({C('cond', r)}=1,0,{C('basem', r)}+{C('vef', r)})",
        "rev": lambda r: (f"=IF(AND(ABS({C('w', r)}-{C('m_adu', r)})<0.01,{C('toyaj', r)}={C('m_toyaj', r)},{C('togaj', r)}={C('m_togaj', r)},"
                          f"{C('zona', r)}={C('m_zona', r)},ABS({C('nec', r)}-{C('m_nec', r)})<0.01),\"OK\",\"DIFIERE\")"),
    }

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Auditoría")
    ws.freeze_panes = f"F{R0}"
    for key, (_, _, _, _, ancho, _) in zip(idx, cols):
        ws.column_dimensions[L(idx[key])].width = ancho
    tipo_font = {"e": AZUL, "f": NEGRO, "m": GRIS}
    tipo_fill = {"e": F_ENT, "f": F_FOR, "m": F_MOT}
    tipo_head = {"e": H_ENT, "f": H_FOR, "m": H_MOT}
    grupo = {"e": "ENTRADAS (valores)", "f": "CÁLCULO (fórmulas)", "m": "LO QUE CALCULÓ EL MOTOR"}

    def hc(texto, tipo, bold=True, wrap=True, size=None):
        c = WriteOnlyCell(ws, value=texto)
        c.font = Font(bold=bold, color="FFFFFF", size=size or 10); c.fill = tipo_head[tipo]
        c.alignment = Alignment(wrap_text=wrap, vertical="center", horizontal="center")
        return c
    # fila 1: grupos, fila 2: titulos, fila 3: descripcion
    ws.append([hc(grupo[c[3]] if (i == 0 or cols[i - 1][3] != c[3]) else None, c[3]) for i, c in enumerate(cols)])
    ws.append([hc(c[1], c[3]) for c in cols])
    desc_cells = []
    for c in cols:
        d = WriteOnlyCell(ws, value=c[2]); d.font = Font(italic=True, size=8, color="5F6368")
        d.alignment = Alignment(wrap_text=True, vertical="top"); d.fill = tipo_fill[c[3]]; desc_cells.append(d)
    ws.append(desc_cells)

    ventas_keys = [f"v{i}" for i in range(n)]
    entradas_map = {"emp": "empresa", "bod": "bodega_id", "ref": "producto_id", "desc": "descripcion", "lab": "laboratorio",
                    "pron": "pronostico", "lt": "lead_time", "frec": "frecuencia", "ud": "unidad_despacho", "cmax": "cota_max", "cmin": "cota_min",
                    "promo": "promo", "vect": "vectorizacion", "inv": "inventario", "tr": "en_transito", "dc": "demanda_calificada", "cond": "condenado", "robj": "regla_objetivo",
                    "m_adu": "adu", "m_toyaj": "toy_aj", "m_togaj": "tog_aj", "m_zona": "zona", "m_nec": "necesidad", "m_cedi": "stock_cedi",
                    "m_asig": "cantidad_sugerida", "m_est": "estado", "m_apr": "cantidad_aprobada"}
    fuentes = {kc: [] for kc in idx}
    for i, fila in enumerate(filas):
        r = R0 + i
        out = []
        for key, titulo, _, tipo, _, fmt in cols:
            if key in F:
                v = F[key](r)
            elif key in ventas_keys:
                v = fila["ventas"][int(key[1:])]
            else:
                v = fila.get(entradas_map[key])
                if isinstance(v, bool): v = int(v)
            cell = WriteOnlyCell(ws, value=v)
            cell.font = tipo_font[tipo]
            if tipo != "f": cell.fill = tipo_fill[tipo]
            if key in ("nec", "m_asig", "rev"): cell.font = Font(bold=True, color=tipo_font[tipo].color)
            if fmt: cell.number_format = fmt
            out.append(cell)
        ws.append(out)
    ws.auto_filter.ref = f"A2:{L(len(cols))}{R0 + max(len(filas), 1) - 1}"

    # ---- hoja Parametros ----
    wp = wb.create_sheet("Parametros")
    wp.column_dimensions["A"].width = 30; wp.column_dimensions["B"].width = 12; wp.column_dimensions["C"].width = 70
    wp.append(["Parámetro", "Valor", "Qué hace (cambia el valor en azul y todas las filas de 'Auditoría' se recalculan)"])
    for clave, desc, defecto in PARAMS:
        v = params.get(clave, defecto)
        c = WriteOnlyCell(wp, value=float(v)); c.font = AZUL; c.fill = F_ENT
        wp.append([clave, c, desc])
    wp.append([]); wp.append(["Tramos de factores (hoja 'Factores' del Excel original)"]); 
    wp.append(["tipo", "descripción", "desde", "hasta", "factor"])
    assert wp is not None
    # los tramos empiezan en la fila TRAMO0 (= off_lt)
    for t in lt_t + va_t:
        cells = [t["tipo"], t["descripcion"]]
        for v in (t["desde"], t["hasta"], t["factor"]):
            c = WriteOnlyCell(wp, value=(None if v is None else float(v))); c.font = AZUL; c.fill = F_ENT; cells.append(c)
        wp.append(cells)
    wp.append([]); wp.append(["Primer tramo: incluye 'desde'. Los demás: mayores que 'desde' y hasta 'hasta' incluido. 'hasta' vacío = sin tope."])

    # ---- hoja Como leerlo ----
    wl = wb.create_sheet("Cómo leerlo")
    wl.column_dimensions["A"].width = 120
    txt = [
        "AUDITORÍA DEL CÁLCULO DDMRP",
        f"Corrida del motor: #{ejecucion['id']} · fecha de cálculo {ejecucion['fecha']} · {len(filas):,} filas" + (f" (de {truncado:,}: filtra más para ver el resto)" if truncado else ""),
        "",
        "Cómo está armada la hoja 'Auditoría' (una fila por bodega × referencia):",
        "  · Azul sobre celeste = ENTRADAS (valores): ventas por mes, pronóstico, lead time, frecuencia, unidad de despacho, cotas, promoción, vectorización, inventario, tránsito, demanda calificada.",
        "  · Negro = CÁLCULO (fórmulas de Excel vivas): el mismo modelo de la hoja 'Modelo' de Descripcion y Formulas DDMRP.xlsx — ADU, zonas, topes, ajuste por cotas/promoción, flujo neto, zona y necesidad.",
        "  · Gris = lo que calculó el motor (valores) para esa misma fila, incluido el reparto del CEDI y el estado de aprobación.",
        "  · La columna REVISIÓN dice OK si la fórmula reproduce el ADU, los topes ajustados, la zona y la necesidad del motor; DIFIERE si algo no coincide.",
        "  · Fila 3 de cada columna explica la fórmula y entre paréntesis la columna equivalente del Excel original.",
        "",
        "De dónde salen las entradas:",
        "  · Inventario, tránsito, demanda calificada, vectorización, lead time, frecuencia, unidad de despacho y condenado: la foto que guardó el motor en esa corrida.",
        "  · Ventas por mes, pronóstico, cotas y promoción: los maestros al momento de descargar. Si los cambiaste después de correr el motor, REVISIÓN puede decir DIFIERE aunque el motor esté bien: recalcula el motor.",
        "",
        "Qué NO es fórmula:",
        "  · El reparto del inventario del CEDI (cuando no alcanza) se calcula en el motor entre todas las tiendas de la referencia; aquí aparece como valor (Stock CEDI repartible y ASIGNADO).",
        "  · NECESIDAD (fórmula) es lo que pide el buffer ANTES de repartir el CEDI. ASIGNADO motor es lo que quedó después.",
        "",
        "Para probar escenarios: cambia los valores azules de la hoja 'Parametros' (regla de disparo, objetivo, redondeo, tramos de factores, multiplicadores de promoción) y mira cómo se mueve la NECESIDAD.",
        "Los rangos de ventas del ADU (últimos meses) y de la ventana quedaron fijos según los parámetros al descargar.",
        "Pequeñas diferencias de 1 unidad pueden venir del redondeo de decimales de Excel frente al del motor.",
    ]
    for t in txt:
        c = WriteOnlyCell(wl, value=t)
        if t.isupper() or t.endswith(":"): c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True); wl.append([c])
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()
