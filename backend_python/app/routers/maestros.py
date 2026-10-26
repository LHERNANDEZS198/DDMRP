"""Maestros editables desde la UI (lead time, cotas, unidad de despacho, vectorizacion, pronostico, promociones,
factores y parametros del modelo). Todo el SQL sale de esta lista blanca: el cliente nunca elige tablas ni columnas."""
import csv, io, re, unicodedata
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from pydantic import BaseModel
from sqlalchemy import text
from sqlalchemy.exc import DBAPIError
from ..db import pg
from .. import bitacora
from ..activos import solo_activos
from ..security import usuario_actual, filtro_sql_permisos, exigir, asignaciones, nivel_efectivo, NIVELES, requiere, permisos

router = APIRouter(prefix="/maestros", tags=["maestros"])

# tipo -> cast SQL
CAST = {"num": "numeric", "int": "smallint", "txt": "text"}
# tipo de las columnas llave que no son texto
TIPO_LLAVE = {"dia": "int"}
# columnas llave cuyo valor numerico de Excel pierde ceros a la izquierda -> se rellenan
ANCHO = {"producto_id": 6, "empresa": 2, "sucursal": 3, "plan_id": 3, "criterio_id": 4}

MAESTROS = {
    "bitacora": dict(
        label="Bitacora: tiendas que se abastecen (nombre, lead time, frecuencia, CEDI y centro de operacion)", tabla="ddmrp.abast_bitacora",
        llave=["empresa", "bodega_id"], editable={"descripcion": "txt", "frecuencia": "num", "leadtime": "num", "cedi": "txt", "id_co": "txt"},
        bodega="bodega_id", producto=None, valida_cedi=True,
        nombre_en_tabla=True,        # el nombre de la tienda vive en la bitacora (fuente de verdad); un archivo con bodega_nombre / tienda / nombre lo guarda aqui
        conservar=("descripcion",)), # una celda de nombre vacia en la carga NO borra el nombre que ya tiene
    "cedis": dict(
        label="CEDIs: bodegas desde donde se abastece a las tiendas", tabla="ddmrp.maestro_cedi",
        llave=["empresa", "bodega_id"], editable={"descripcion": "txt"}, bodega=None, producto=None, global_=True, empresa_obligatoria=True),
    "control-referencias": dict(
        label="Cotas maxima / minima por bodega y referencia", tabla="ddmrp.abast_control_referencias",
        llave=["empresa", "bodega_id", "producto_id"], editable={"cota_maxima": "num", "cota_minima": "num"},
        bodega="bodega_id", producto="producto_id"),
    "vectorizacion": dict(
        label="Vectorizacion (derrame); NxP = condenar item en esa tienda", tabla="ddmrp.abast_vectorizacion",
        llave=["empresa", "bodega_id", "producto_id"], editable={"vectorizacion": "num"},
        bodega="bodega_id", producto="producto_id", nxp=True),   # NxP en vectorizacion = condenado (item x tienda): va al maestro de condenados
    "condenados": dict(
        label="Condenados: items que no se despachan (todas las tiendas o solo una)", tabla="ddmrp.item_condenado",
        llave=["empresa", "bodega_id", "producto_id"], editable={"motivo": "txt"}, bodega="bodega_id", producto="producto_id", todas="*"),
    "leadtime-cedi": dict(
        label="Lead time del CEDI (dias): general (*) o por item", tabla="ddmrp.abast_leadtime_cedi",
        llave=["empresa", "producto_id"], editable={"leadtime": "num"},
        bodega=None, producto="producto_id", global_=True, grupo="compras"),
    "proveedores": dict(
        label="Proveedores: lead time, correos, estado y pedido minimo", tabla="ddmrp.proveedor",
        llave=["empresa", "proveedor_id", "sucursal"],
        editable={"nombre": "txt", "leadtime": "num", "email_para": "txt", "email_cc": "txt", "activo": "int", "redondeo": "txt", "frecuencia": "num", "pedido_minimo": "num",
                  "contacto": "txt", "telefono": "txt", "notas": "txt"},
        bodega=None, producto=None, global_=True, grupo="compras",
        buscar=("proveedor_id", "nombre"),      # el cuadro de busqueda mira el NIT Y el nombre
        emails=("email_para", "email_cc"),      # varios correos separados por ; (o coma): se validan y se normalizan
        booleanas=("activo",),                  # 1 / 0 / si / no
        valores={"redondeo": ("arriba", "cercano", "abajo")},   # como se redondea la cantidad al multiplo de pedido
        conservar=("activo", "redondeo", "frecuencia")),                 # celda vacia en la carga = no cambia
    "config-compras": dict(
        label="Configuracion de compras: pedido minimo y poca historia", tabla="ddmrp.compras_config",
        llave=["clave"], editable={"valor": "num", "detalle": "txt"},
        bodega=None, producto=None, global_=True, grupo="compras",
        solo_existentes=True, sin_reemplazo=True, sin_eliminar=True, conservar=("valor", "detalle")),
    "criterios-comprables": dict(
        label="Criterios comprables: que categorias se compran en cada empresa", tabla="ddmrp.categoria_compra",
        llave=["empresa", "plan_id", "criterio_id"], editable={"nombre": "txt", "se_compra": "int"},
        bodega=None, producto=None, global_=True, grupo="compras",
        booleanas=("se_compra",), conservar=("nombre", "se_compra")),
    "items-compras": dict(
        label="Items de compras: cotizacion, unidad de pedido y unidad de empaque", tabla="ddmrp.item_compra",
        llave=["empresa", "producto_id"], editable={"cotizacion": "num", "unidad_pedido": "num", "unidad_empaque": "num", "moq": "num"},
        bodega=None, producto="producto_id", global_=True, grupo="compras",
        solo_existentes=True,                   # los items vienen de Siesa: la carga solo actualiza, nunca crea
        sin_reemplazo=True, sin_eliminar=True,  # y aqui no se borra nada (ni filas sueltas, ni vaciar, ni reemplazar)
        conservar=("cotizacion", "unidad_pedido", "unidad_empaque", "moq"),   # celda vacia = no cambia
        lectura=[("laboratorio", "Laboratorio"), ("categoria", "Categoría"), ("linea", "Línea"), ("exhibicion", "Exhibición"), ("costo", "Costo promedio"),
                 ("cotizacion_siesa", "Cotización Siesa"), ("cotizacion_vigencia", "Vigente desde"), ("ultimo_costo", "Último costo de compra"),
                 ("ultimo_costo_fecha", "Fecha última compra"), ("costo_compra", "Costo a usar"), ("origen_costo", "Origen del costo"),
                 ("iva_tasa", "IVA % (estimado)"), ("proveedor", "Proveedor (NIT)"), ("proveedor_nombre", "Nombre del proveedor"),
                 ("origen_proveedor", "Origen del proveedor"), ("proveedor_item", "Proveedor en el ítem (NIT)"), ("abastecimiento", "Abastecimiento")]),
    "pronostico": dict(
        label="Pronostico (unidades/mes)", tabla="ddmrp.pronostico",
        llave=["empresa", "bodega_id", "producto_id"], editable={"pronostico": "num"},
        bodega="bodega_id", producto="producto_id"),
    "unidad-despacho": dict(
        label="Unidad de despacho por referencia", tabla="ddmrp.abast_unidad_despacho",
        llave=["empresa", "producto_id"], editable={"unidad_despacho": "num"},
        bodega=None, producto="producto_id", global_=True),
    "reserva-cedi": dict(
        label="Reserva del CEDI: % del inventario que se guarda por item (no se despacha)", tabla="ddmrp.reserva_cedi",
        llave=["producto_id"], editable={"porcentaje": "num"}, bodega=None, producto="producto_id", global_=True, fraccion=True),
    "promocion": dict(
        label="Promociones (indicativo 0/1/2)", tabla="ddmrp.promocion",
        llave=["empresa", "producto_id"], editable={"indicativo": "int"},
        bodega=None, producto="producto_id", global_=True),
    "factores": dict(
        label="Factores de lead time y variabilidad", tabla="ddmrp.factor_tramo",
        llave=["tipo", "descripcion"], editable={"factor": "num", "desde": "num", "hasta": "num"},
        bodega=None, producto=None, global_=True),
    "parametros": dict(
        label="Parametros del modelo", tabla="ddmrp.parametro",
        llave=["clave"], editable={"valor": "num"}, bodega=None, producto=None, global_=True,
        solo_existentes=True, sin_reemplazo=True, oculto=True),   # oculto: las mismas claves se editan en Reglas de pedido (con validacion y nombres claros)
    "dias-bodega": dict(
        label="DIA de abastecimiento por tienda", tabla="ddmrp.bodega_dia",
        llave=["empresa", "bodega_id", "dia"], editable={}, bodega="bodega_id", producto=None),
    "items": dict(
        label="Items: descripcion y laboratorio (los criterios vienen de Siesa)", tabla="ddmrp.maestro_item",
        llave=["producto_id"], editable={"descripcion": "txt", "categoria": "txt"},
        bodega=None, producto=None, global_=True, defecto={"descripcion": "producto_id"}),
    "items-rec": dict(
        label="Items REC: cada item REC y su original (el SKU sin el primer digito)", tabla="ddmrp.item_rec",
        llave=["empresa", "rec_id"], editable={"original_id": "txt"}, bodega=None, producto=None, global_=True),
    "criterios-filtros": dict(
        label="Criterios de items que salen en los filtros (1 = sale · 0 = no sale)", tabla="ddmrp.criterio_plan",
        llave=["plan_id"], editable={"en_filtros": "int"}, bodega=None, producto=None, global_=True,
        solo_existentes=True, sin_reemplazo=True),
    "reposicion-laboratorio": dict(
        label="Hasta donde se repone, por laboratorio (1 TOG aj · 2 TOY aj · 3 TOY aj + mitad de la zona verde)", tabla="ddmrp.reposicion_laboratorio",
        llave=["laboratorio"], editable={"regla": "int"}, bodega=None, producto=None, global_=True),
    "reposicion-item": dict(
        label="Hasta donde se repone, por item (1 TOG aj · 2 TOY aj · 3 TOY aj + mitad de la zona verde)", tabla="ddmrp.reposicion_item",
        llave=["producto_id"], editable={"regla": "int"}, bodega=None, producto="producto_id", global_=True),
    "laboratorio-sin-envio": dict(
        label="Laboratorios que NO se envian (no se sugieren, no compiten por el CEDI ni salen en el plano)", tabla="ddmrp.laboratorio_sin_envio",
        llave=["laboratorio"], editable={"motivo": "txt"}, bodega=None, producto=None, global_=True),
    "demanda-calificada": dict(
        label="Demanda calificada (pedidos abiertos / picos)", tabla="ddmrp.demanda_calificada",
        llave=["empresa", "bodega_id", "producto_id"], editable={"cantidad": "num"},
        bodega="bodega_id", producto="producto_id"),
    "ventas-mensual": dict(
        label="Ventas por mes (unidades)", tabla="ddmrp.ventas_mensual",
        llave=["empresa", "mes", "bodega_id", "producto_id"], editable={"cantidad": "num"},
        bodega="bodega_id", producto="producto_id"),
}


SIN_VACIAR = {"items", "factores", "parametros", "items-compras"}   # se llenan desde Siesa o son del modelo: no se vacian desde la pantalla
OPS = {"eq": "=", "ne": "<>", "gt": ">", "lt": "<", "ge": ">=", "le": "<="}

AYUDA = {
    "bitacora": "Las tiendas que se abastecen y su nombre: ESTA es la fuente de verdad de las tiendas (el nombre que escribas aquí es el que se ve en todo el sistema; Siesa solo se usa si no hay nombre). Solo las que tú metas aquí (nada se agrega solo). Por tienda: nombre, lead time (días que tarda el producto en llegar a la tienda), frecuencia (cada cuántos días se abastece), "
                "CEDI que la abastece (código de bodega del maestro CEDIs, ej. PT001) y centro de operación para el archivo plano de Siesa (si lo dejas vacío, el ETL lo completa desde Siesa y si no hay, se deduce con la regla del ejemplo). "
                "El lead time × la venta diaria da la zona amarilla y la frecuencia × la venta diaria da la verde; además el lead time define el factor de la zona roja (maestro Factores). "
                "Una tienda que no esté aquí no entra al cálculo del motor.",
    "control-referencias": "Cota máxima y mínima por tienda y referencia, como fracción o % (20% = 0.2). Ajustan los topes: TOY ajustado = TOG × cota máx + TOR + ZA, "
                           "TOG ajustado = TOY ajustado + ZV + TOG × cota mín. Con cotas en 0 los topes ajustados quedan iguales a los calculados. No se usan cuando la venta diaria es baja (ADU entre 1 y 4) y no hay promoción.",
    "vectorizacion": "Unidades extra (derrame) que se suman a lo sugerido de esa tienda y referencia, p. ej. producto nuevo o campaña. Solo cuenta si la regla 'Sumar vectorización' está en Sí. "
                     "Escribir NxP condena el item en esa tienda (pasa al maestro Condenados).",
    "condenados": "Items que NO se despachan: el buffer se calcula y se ve, pero la cantidad sugerida queda en 0. Tienda * = el item queda condenado en todas las tiendas.",
    "unidad-despacho": "Múltiplo en que se despacha cada referencia (caja o display). Lo sugerido se redondea a este múltiplo (según la regla de redondeo) y también es el mínimo de la zona verde.",
    "reserva-cedi": "Porcentaje del inventario del CEDI que se guarda por item: no se reparte a las tiendas. Solo importa cuando el CEDI no alcanza para todas.",
    "promocion": "Indicativo de promoción por referencia. 0 = sin promoción. 1 = promoción fuerte: las cotas se agrandan (máxima × 1,2 y mínima × 1,1, parámetros del modelo), así el buffer crece y se sugiere más. "
                 "2 = promoción moderada: se aplican las cotas normales aunque la venta diaria sea baja (la mínima × 1,1 solo en tiendas GN). "
                 "Excepción: la promoción NO aplica en las tiendas de confianza (código de bodega que empieza por TC) ni en los domicilios (DM): ahí el indicativo se ignora (se toma como 0) y el buffer se calcula sin promoción. "
                 "Los prefijos son los dos primeros caracteres del código de bodega: GN Global Nutrition, SV Sol Verde, NS Natures Stores, MN Melonn (tiendas reales); TC tiendas de confianza y DM domicilios (sin promoción).",
    "factores": "Tablas de tramos del modelo. LT: factor de la zona roja base según el lead time. VAR: factor de la zona roja de seguridad según la variabilidad de la demanda. Cambian los buffers de todos los productos.",
    "parametros": "Parámetros del modelo: multiplicadores de cota en promoción y rango de venta diaria baja (ADU). Solo se editan los que ya existen. Las reglas de pedido están en la pantalla Reglas de pedido.",
    "dias-bodega": "DÍA(S) de abastecimiento de cada tienda. Sirve para armar la corrida del día en la Matriz de despacho y, si lo eliges en Reglas de pedido, para repartir el inventario del CEDI solo entre las tiendas de esos días.",
    "laboratorio-sin-envio": "Laboratorios (el código de la columna Laboratorio del maestro de items, ej. 1781 o 181) que no se envían: sus items no se sugieren en la Matriz de despacho, "
                             "no compiten por el stock del CEDI y no salen en el plano de Siesa. Se aplica en la próxima corrida del motor (la matriz y el plano lo respetan de inmediato).",
    "items": "Descripción y laboratorio de cada item; los demás criterios vienen de Siesa con 'Traer items'. El laboratorio define los permisos por categoría y los filtros del Control.",
    "items-rec": "Items cuya descripción empieza por REC (ej. 2190340 REC CALENDULA...): cuentan como su original (190340, el SKU sin el primer dígito). Su inventario, tránsito y ventas se suman al original y el REC sale del cálculo. "
                 "El ETL detecta los pares solo (origen auto) cuando el original existe en Siesa; aquí puedes agregar o corregir uno a mano (origen manual, no lo pisa el ETL). Se aplica con la regla 'Items REC se suman a su original'. Para despachar, el stock debe estar físicamente bajo el SKU original en Siesa.",
    "cedis": "CEDIs (bodegas desde donde se despacha a las tiendas), uno por empresa y código de bodega, ej. empresa 05 PT001. La bitácora dice qué CEDI abastece a cada tienda. "
             "Hoy el motor reparte el inventario de un solo CEDI por empresa (el primero de esta lista); la columna CEDI de la bitácora deja listo el modelo para cuando haya más de uno.",
    "proveedores": "Los proveedores llegan solos del maestro de items de Siesa (cada vez que se traen los items entran los nuevos, con su NIT y nombre) y aqui compras completa lo suyo: lead time en dias "
                   "(prioridad al calcular el CEDI: el del item en «Lead time del CEDI», luego el del proveedor, luego el general), frecuencia de compra en dias (cada cuanto se le compra; define el tamaño de la zona verde del CEDI: a mayor frecuencia, pedidos mas grandes; vacio = el general de «Configuracion de compras», 15 dias), redondeo (como se lleva la cantidad a pedir al multiplo de pedido del item: arriba [por defecto], cercano o abajo), correos (varios separados por ;), activo (1/0 o si/no; vacio = activo), "
                   "pedido minimo (valor), contacto, telefono y notas. Lo que ya escribiste no se pisa al sincronizar. La llave es empresa + NIT + sucursal.",
    "items-compras": "Los items vienen de Siesa (no se crean ni se borran aqui) y casi todas las columnas son de solo lectura. "
                     "PROVEEDOR: es el de la COTIZACION vigente del item (si hay cotizaciones de varios proveedores, la de mayor fecha de activacion); solo si no hay cotizacion se usa el que trae el item en Siesa. "
                     "«Origen del proveedor» dice de donde salio y «Proveedor en el item» muestra el del item (suele estar desactualizado, sobre todo en maquila). "
                     "COSTO A USAR: la cotizacion manual; si no hay, la cotizacion de Siesa (costo bruto con el descuento del primer rango de cantidad); si tampoco, el ULTIMO COSTO DE COMPRA (precio de la ultima OC aprobada, parcial o cumplida); «Origen del costo» dice cual. "
                     "El IVA es una tarifa ESTIMADA: Siesa lo calcula al crear la OC. "
                     "CATEGORIA = criterio del plan 001 de Siesa (0001 Productos Funat, 0002 Productos de terceros, 0003 Maquila Funat). ABASTECIMIENTO: Compra, Produccion propia (la categoria no se compra en esa empresa: maestro Criterios comprables) o Sin proveedor. "
                     "Compras carga a mano la cotizacion (solo si quiere pisar las demas), la unidad de pedido (multiplo en que se pide, por defecto 1), la unidad de empaque (unidades por caja, por defecto 1) y el MOQ (cantidad minima de compra en unidades, por defecto 1: si el pedido queda por debajo, se sube a ese minimo al multiplo de pedido). "
                     "Por Excel: baja el archivo, llena esas columnas y subelo; celda vacia = no cambia. Solo se ven los items activos de su empresa.",
    "config-compras": "Ajustes generales del informe Necesidades de compra. pedido_minimo_bloquea: 0 = los proveedores cuyo total queda bajo su pedido minimo se muestran con una alerta (no se bloquean); 1 = se bloquean y no se muestran. "
                      "poca_historia_meses: un item con menos de estos meses con venta (de los ultimos 12) se marca «poca historia»: solo alerta, no bloquea. Solo se edita el valor.",
    "criterios-comprables": "Que categorias (criterio del plan 001 de Siesa) se COMPRAN en cada empresa: 1 = se compra, 0 = produccion propia (no genera OC). Funat (05): terceros y maquila se compran y Productos Funat no; Stores (21): todo se compra. "
                            "Una categoria que no aparece aqui se trata como compra. El codigo del criterio va como texto de 4 digitos (ej. 0003).",
    "leadtime-cedi": "Dias que tarda en llegar un item al CEDI (compra o maquila; la produccion propia no se modela). La fila con producto * es el general de la empresa (8 dias); "
                     "una fila con un item lo sobrescribe y, sin fila de item, manda el lead time del proveedor del item (maestro Proveedores). Alimenta el buffer del CEDI (Informes > Ver solo CEDI): se aplica en la siguiente corrida del motor.",
    "demanda-calificada": "Pedidos abiertos o picos conocidos por tienda y referencia. Se restan del flujo neto cuando la regla 'Restar demanda calificada' está en Sí.",
    "pronostico": "Venta mensual esperada (unidades/mes) por tienda y referencia. Sirve sobre todo para tiendas, referencias o laboratorios NUEVOS: si la combinacion tiene menos meses de venta que la regla 'Pronostico directo en casos nuevos' (por defecto 3), el motor usa este valor tal cual, aunque no haya historial. Con mas historial el pronostico solo ajusta: el motor lo lleva entre 1 y 2 veces el ADU de referencia (nunca baja el historial). Para campanas o promociones usa Promocion y las cotas.",
    "ventas-mensual": "Unidades vendidas por mes, tienda y referencia: base de la venta diaria (ADU) y de la variabilidad. Se traen de Siesa con 'Traer ventas por mes' o se cargan por Excel.",
    "criterios-filtros": "Planes de criterios de items traídos de Siesa (CATEGORIA, LINEA...). En_filtros = 1 hace que el plan salga como filtro y como 'Agrupar por' en Clasificación, Control y Lanzamientos; 0 lo oculta. "
                         "Un plan nuevo que llega de Siesa sale en 1 (visible). No afecta las columnas del Excel de Clasificación ni la lista de Items.",
    "reposicion-laboratorio": "Hasta dónde se repone cuando se pide, para todos los items de un laboratorio (código del laboratorio): 1 = TOG ajustado (tope verde), 2 = TOY ajustado (tope amarillo), "
                              "3 = tope amarillo + la mitad de la zona verde ajustada. Sobrescribe la regla general de Reglas de pedido; la del item gana sobre la del laboratorio. Aplica en la siguiente corrida del motor.",
    "reposicion-item": "Hasta dónde se repone cuando se pide, para un item en todas las tiendas: 1 = TOG ajustado (tope verde), 2 = TOY ajustado (tope amarillo), 3 = tope amarillo + la mitad de la zona verde ajustada. "
                       "Es la prioridad más alta: gana sobre la del laboratorio y sobre la regla general. Aplica en la siguiente corrida del motor.",
}


def _def(nombre: str) -> dict:
    d = MAESTROS.get(nombre)
    if not d:
        raise HTTPException(404, "Maestro desconocido")
    return d


# Cada maestro pertenece a un grupo: "ddmrp" (abastecimiento a tiendas, el de siempre) o "compras" (lo que usa el area de compras).
# Cada grupo tiene su propio permiso de ver y de cargar; el administrador tiene todos.
PERM_GRUPO = {"ddmrp": (("ver_maestros", "cargar_maestros"), ("cargar_maestros",)),
              "compras": (("ver_compras", "cargar_maestros_compras"), ("cargar_maestros_compras",))}


def _grupo(d: dict) -> str:
    return d.get("grupo", "ddmrp")


def _def_para(nombre: str, user, escritura: bool = False) -> dict:
    """El maestro, solo si el perfil del usuario tiene permiso sobre SU grupo (ver o, con escritura, cargar)."""
    d = _def(nombre)
    ver, esc = PERM_GRUPO[_grupo(d)]
    if not permisos(user["id"]).intersection(esc if escritura else ver):
        raise HTTPException(403, f"Tu perfil no tiene permiso sobre los maestros de {_grupo(d)}")
    return d


def _crit_plan(patron: str) -> str:
    """Nombre del criterio del item (de Siesa) en el plan cuyo nombre empieza por `patron` (LIKE, sin mayusculas)."""
    return ("(SELECT cr.nombre FROM ddmrp.item_criterio ic JOIN ddmrp.criterio_plan cp ON cp.plan_id = ic.plan_id "
            "LEFT JOIN ddmrp.criterio cr ON cr.plan_id = ic.plan_id AND cr.criterio_id = ic.criterio_id "
            f"WHERE ic.producto_id = t.producto_id AND cp.nombre ILIKE '{patron}' LIMIT 1)")


# Categoria de compras = criterio del plan 001 de Siesa (0001 Productos Funat, 0002 Productos de terceros, 0003 Maquila Funat).
_CATID_SQL = "(SELECT ic.criterio_id FROM ddmrp.item_criterio ic WHERE ic.producto_id = t.producto_id AND ic.plan_id = '001' LIMIT 1)"
_CAT_SQL = ("(SELECT CASE WHEN cr.nombre IS NULL THEN ic.criterio_id ELSE ic.criterio_id || ' · ' || cr.nombre END FROM ddmrp.item_criterio ic "
            "LEFT JOIN ddmrp.criterio cr ON cr.plan_id = ic.plan_id AND cr.criterio_id = ic.criterio_id WHERE ic.producto_id = t.producto_id AND ic.plan_id = '001' LIMIT 1)")
# Abastecimiento: Produccion propia (la categoria no se compra en esa empresa, maestro «Criterios comprables»), Sin proveedor (ni cotizacion ni proveedor en el item) o Compra.
_ABAS_SQL = (f"CASE WHEN COALESCE((SELECT cc.se_compra FROM ddmrp.categoria_compra cc WHERE cc.empresa = t.empresa AND cc.plan_id = '001' AND cc.criterio_id = {_CATID_SQL}), 1) = 0 "
             "THEN 'Producción propia' WHEN t.proveedor_id IS NULL THEN 'Sin proveedor' ELSE 'Compra' END")


def _filtro_lectura(categoria: str | None, abastecimiento: str | None) -> tuple[str, dict]:
    """Filtros de las columnas calculadas de items-compras (listas separadas por coma; categoria = codigos de criterio: 0001, 0002, 0003)."""
    w, p = "", {}
    for clave, expr, vals in (("categoria", _CATID_SQL, categoria), ("abastecimiento", _ABAS_SQL, abastecimiento)):
        lista = [x.strip() for x in (vals or "").split(",") if x.strip()]
        if lista:
            w += f" AND ({expr}) = ANY(:fl_{clave})"; p[f"fl_{clave}"] = lista
    return w, p


def _lectura_compras() -> tuple[str, str, list[str]]:
    """(joins, columnas, encabezados) de solo lectura del maestro items-compras. El nombre sale del maestro de items (alias mi, ya unido).
    LINEA y EXHIBICION son planes de criterios de Siesa: se buscan por el nombre del plan."""
    join = (" LEFT JOIN ddmrp.laboratorio lb ON lb.codigo = mi.categoria LEFT JOIN ddmrp.costo_item ci ON ci.producto_id = t.producto_id"
            " LEFT JOIN ddmrp.proveedor pv ON pv.empresa = t.empresa AND pv.proveedor_id = t.proveedor_id AND pv.sucursal = t.sucursal")
    sel = (", CASE WHEN lb.nombre IS NULL THEN mi.categoria ELSE mi.categoria || ' · ' || lb.nombre END AS laboratorio, " + _CAT_SQL + " AS categoria, "
           + _crit_plan("L_NEA%") + " AS linea, " + _crit_plan("EXHIBIC%") + " AS exhibicion, "
           "ci.costo::float8 AS costo, t.cotizacion_siesa::float8 AS cotizacion_siesa, t.cotizacion_fecha AS cotizacion_vigencia, "
           "t.ultimo_costo::float8 AS ultimo_costo, t.ultimo_costo_fecha AS ultimo_costo_fecha, "
           "COALESCE(t.cotizacion, t.cotizacion_siesa, t.ultimo_costo)::float8 AS costo_compra, "
           "CASE WHEN t.cotizacion IS NOT NULL THEN 'Manual' WHEN t.cotizacion_siesa IS NOT NULL THEN 'Cotización Siesa' WHEN t.ultimo_costo IS NOT NULL THEN 'Última compra' ELSE 'Sin costo' END AS origen_costo, "
           "t.iva_tasa::float8 AS iva_tasa, t.proveedor_id AS proveedor, pv.nombre AS proveedor_nombre, "
           "CASE t.proveedor_origen WHEN 'cotizacion' THEN 'Cotización' WHEN 'item' THEN 'Ítem' END AS origen_proveedor, t.proveedor_item AS proveedor_item, " + _ABAS_SQL + " AS abastecimiento")
    return join, sel, ["descripcion", "laboratorio", "categoria", "linea", "exhibicion", "costo", "cotizacion_siesa", "cotizacion_vigencia", "ultimo_costo", "ultimo_costo_fecha", "costo_compra", "origen_costo", "iva_tasa", "proveedor", "proveedor_nombre", "origen_proveedor", "proveedor_item", "abastecimiento"]


def _es_admin(user) -> bool:
    return any(a["es_admin"] for a in asignaciones(user["id"]))


def _puede_global(user) -> bool:
    """Tablas sin bodega: admin, o edicion/aprobacion sobre todas las bodegas y categorias."""
    return any(a["es_admin"] or (a["bodega_id"] == "*" and a["categoria"] == "*" and NIVELES[a["nivel"]] >= 2)
               for a in asignaciones(user["id"]))


def _autorizar(user, d: dict, llave: dict, conn):
    if d.get("bodega"):
        cat = None
        if d.get("producto"):
            cat = conn.execute(text("SELECT categoria FROM ddmrp.maestro_item WHERE producto_id=:p"),
                               {"p": llave[d["producto"]]}).scalar()
        exigir(user, llave[d["bodega"]], cat, "edicion")
    elif d.get("global_"):
        if not _puede_global(user):
            raise HTTPException(403, "Requiere permiso de edicion sobre todas las bodegas")
    elif not _es_admin(user):
        raise HTTPException(403, "Solo administradores")


@router.get("")
def listar_maestros(user=Depends(requiere("ver_maestros", "cargar_maestros", "ver_compras", "cargar_maestros_compras"))):
    mios = permisos(user["id"])
    return [{"clave": k, "grupo": _grupo(v), "label": v["label"], "llave": v["llave"], "editable": list(v["editable"]),
             "tipos": v["editable"], "sin_reemplazo": bool(v.get("sin_reemplazo")), "todas": bool(v.get("todas")),
             "solo_existentes": bool(v.get("solo_existentes")), "lectura": [{"k": a, "t": b} for a, b in v.get("lectura", [])],
             "ayuda": AYUDA.get(k, ""),
             "puede_vaciar": not v.get("sin_reemplazo") and k not in SIN_VACIAR,
             "filtros": {"tienda": bool(v.get("bodega")), "item": bool(v.get("producto")) or k == "items",
                         "lab": bool(v.get("producto")) or k in ("items", "reposicion-laboratorio", "laboratorio-sin-envio"), "texto": True,
                         "empresa": k == "items",                                   # items: la empresa es obligatoria (cada base tiene su catalogo)
                         "empresa_opc": k != "items" and "empresa" in v["llave"],   # los demas: opcional (vacio = las dos)
                         "mes": k == "ventas-mensual"}}
            for k, v in MAESTROS.items() if not v.get("oculto") and mios.intersection(PERM_GRUPO[_grupo(v)][0])]


@router.get("/_opciones")
def opciones(user=Depends(requiere("ver_maestros", "cargar_maestros", "ver_compras", "cargar_maestros_compras"))):
    """Listas para los filtros de la pantalla: tiendas (con nombre y empresa) y laboratorios. Va antes de /{nombre}."""
    w, p = filtro_sql_permisos(user, "mb.bodega_id", "'*'")
    with pg.connect() as c:
        tiendas = [dict(r) for r in c.execute(text(
            f"SELECT mb.bodega_id, string_agg(DISTINCT mb.descripcion, ' / ') AS nombre, string_agg(DISTINCT mb.empresa, ',') AS empresa FROM ddmrp.maestro_bodega mb WHERE {w} GROUP BY mb.bodega_id ORDER BY 1"), p).mappings()]
        labs = [dict(r) for r in c.execute(text(
            "SELECT l.codigo, l.nombre FROM ddmrp.laboratorio l WHERE EXISTS (SELECT 1 FROM ddmrp.maestro_item m WHERE m.categoria = l.codigo) ORDER BY 1")).mappings()]
    return {"tiendas": tiendas, "laboratorios": labs}


def _filtros_maestro(nombre: str, d: dict, bodega_id: str | None, q: str | None, lab: str | None,
                     empresa: str | None = None, mes: str | None = None, estado: str | None = None,
                     col: str | None = None, op: str | None = None, val: str | None = None) -> tuple[str, dict]:
    """Filtros de pantalla y de descarga. Supone el alias t (tabla) y mi (maestro_item) cuando el maestro tiene producto."""
    w, p = "", {}
    bod, prod = d.get("bodega"), d.get("producto")
    if bod and bodega_id:
        w += f" AND t.{bod} = ANY(:fb)"; p["fb"] = [x.strip() for x in bodega_id.split(",") if x.strip()]
    if q and q.strip():
        p["q"] = f"%{q.strip()}%"
        if prod:
            w += f" AND (t.{prod} ILIKE :q OR mi.descripcion ILIKE :q)"
        elif nombre == "items":
            w += " AND (t.producto_id ILIKE :q OR t.descripcion ILIKE :q)"
        elif d.get("buscar"):   # el maestro declara en que columnas se busca (p.ej. proveedores: NIT y nombre)
            w += " AND (" + " OR ".join(f"CAST(t.{c} AS text) ILIKE :q" for c in d["buscar"]) + ")"
        else:
            clave = next((k for k in d["llave"] if k != "empresa"), d["llave"][0])
            w += f" AND CAST(t.{clave} AS text) ILIKE :q"
    labs = [x.strip() for x in (lab or "").split(",") if x.strip()]
    if labs:   # uno o varios laboratorios (codigos)
        p["lab"] = labs
        if prod:
            w += " AND mi.categoria = ANY(:lab)"
        elif nombre == "items":
            w += " AND t.categoria = ANY(:lab)"
        elif nombre in ("reposicion-laboratorio", "laboratorio-sin-envio"):
            w += " AND t.laboratorio = ANY(:lab)"
    if nombre == "items":   # items de una sola empresa (las dos bases tienen catalogos distintos)
        if empresa not in ("05", "21"):
            raise HTTPException(400, "Elige la empresa (05 o 21) para ver los items")
        w += " AND EXISTS (SELECT 1 FROM ddmrp.item_empresa ie WHERE ie.producto_id = t.producto_id AND ie.empresa = :emp)"; p["emp"] = empresa
        est = (estado or "activo").lower()   # estado en Siesa (ETL): por defecto solo los activos; sin dato de estado = activo
        if est not in ("activo", "inactivo", "bloqueado", "todos"):
            raise HTTPException(400, "estado debe ser activo, inactivo, bloqueado o todos")
        if est == "activo":
            w += " AND NOT EXISTS (SELECT 1 FROM ddmrp.item_estado ix WHERE ix.empresa = :emp AND ix.producto_id = t.producto_id AND ix.estado <> 1)"
        elif est != "todos":
            w += " AND EXISTS (SELECT 1 FROM ddmrp.item_estado ix WHERE ix.empresa = :emp AND ix.producto_id = t.producto_id AND ix.estado = :est)"
            p["est"] = 0 if est == "inactivo" else 2
    if nombre == "ventas-mensual" and mes and mes.strip():   # filtro de mes en el servidor (no solo sobre las filas ya cargadas)
        p["mes"] = mes.strip().replace("-", "")[:6]
        w += " AND t.mes = :mes"
    if nombre != "items" and "empresa" in d["llave"] and empresa in ("05", "21"):   # maestros con empresa en la llave: filtro opcional
        w += " AND t.empresa = :emp"; p["emp"] = empresa
    if col:   # condicion sobre una columna (llave o editable), resuelta en la base: sirve igual con 400.000 filas
        tipos = {**{k: TIPO_LLAVE.get(k, "txt") for k in d["llave"]}, **d["editable"]}
        if col not in tipos:
            raise HTTPException(400, f"Columna no valida para filtrar: {col}")
        tipo, op = tipos[col], (op or "eq")
        if op == "vacio":
            w += f" AND (t.{col} IS NULL" + (f" OR CAST(t.{col} AS text) = ''" if tipo == "txt" else "") + ")"
        elif op == "lleno":
            w += f" AND t.{col} IS NOT NULL" + (f" AND CAST(t.{col} AS text) <> ''" if tipo == "txt" else "")
        elif val is None or not str(val).strip():
            pass                                    # condicion a medias (sin valor): no filtra
        elif op == "contiene":
            w += f" AND CAST(t.{col} AS text) ILIKE :cv"; p["cv"] = f"%{str(val).strip()}%"
        elif op in OPS:
            if tipo in ("num", "int"):
                try:
                    p["cv"] = float(str(val).strip().replace(",", "."))
                except ValueError:
                    raise HTTPException(400, f"{col} es numerica: {val!r} no es un numero")
                w += f" AND t.{col} {OPS[op]} CAST(:cv AS numeric)"
            else:
                p["cv"] = str(val).strip()
                w += f" AND CAST(t.{col} AS text) {OPS[op]} :cv"
        else:
            raise HTTPException(400, f"Operador no valido: {op}")
    return w, p


@router.get("/{nombre}")
def leer(nombre: str, response: Response, bodega_id: str | None = None, producto_id: str | None = None, q: str | None = None,
         lab: str | None = None, empresa: str | None = None, mes: str | None = None, estado: str | None = None, limit: int = 20000,
         col: str | None = None, op: str | None = None, val: str | None = None, orden: str | None = None, desc: bool = False,
         categoria: str | None = None, abastecimiento: str | None = None,
         user=Depends(requiere("ver_maestros", "cargar_maestros", "ver_compras", "cargar_maestros_compras"))):
    d = _def_para(nombre, user, False)
    cols = d["llave"] + [c for c in d["editable"] if c not in d["llave"]]
    sel = ", ".join(f"t.{c}" for c in cols)
    join, w, p = "", "TRUE", {}
    if d.get("bodega"):
        join = f"LEFT JOIN ddmrp.maestro_item mi ON mi.producto_id = t.{d['producto']}" if d.get("producto") else ""
        w, p = filtro_sql_permisos(user, f"t.{d['bodega']}", "mi.categoria" if d.get("producto") else "'*'")
        sel += ", mi.descripcion" if d.get("producto") else ""
        join += f" LEFT JOIN ddmrp.maestro_bodega mb ON mb.bodega_id = t.{d['bodega']}" + (" AND mb.empresa = t.empresa" if "empresa" in d["llave"] else "")
        if not d.get("nombre_en_tabla"):   # bitacora: el nombre es SU columna descripcion, no el del catalogo
            sel += ", mb.descripcion AS bodega_nombre"
    elif d.get("producto"):
        join = f"LEFT JOIN ddmrp.maestro_item mi ON mi.producto_id = t.{d['producto']}"
        sel += ", mi.descripcion"
    if nombre == "criterios-filtros":
        sel += ", t.nombre AS descripcion"
    if d.get("lectura"):   # items-compras: solo items vigentes (en el maestro de su empresa) y activos, con sus columnas de solo lectura
        j, s_, _ = _lectura_compras(); join += j; sel += s_
        w += " AND EXISTS (SELECT 1 FROM ddmrp.item_empresa ie WHERE ie.empresa = t.empresa AND ie.producto_id = t.producto_id) AND " + solo_activos("t")
        fw, fp = _filtro_lectura(categoria, abastecimiento); w += fw; p.update(fp)
    if producto_id and d.get("producto"):
        w += f" AND t.{d['producto']} = :pr"; p["pr"] = producto_id
    wf, pf = _filtros_maestro(nombre, d, bodega_id, q, lab, empresa, mes, estado, col, op, val)
    w, p = w + wf, {**p, **pf}
    if nombre == "items":   # costo promedio por instalacion (lo trae el ETL) y estado en Siesa: solo lectura
        join += " LEFT JOIN ddmrp.costo_item ci ON ci.producto_id = t.producto_id"
        sel += ", ci.costo::float8 AS costo, " + "(SELECT CASE ix2.estado WHEN 1 THEN \'Activo\' WHEN 0 THEN \'Inactivo\' WHEN 2 THEN \'Bloqueado\' END FROM ddmrp.item_estado ix2 WHERE ix2.empresa = :emp AND ix2.producto_id = t.producto_id) AS estado"
    # orden en el servidor (la pantalla solo trae un pedazo cuando el maestro es grande): solo por columnas de la tabla
    propias = set(d["llave"]) | set(d["editable"])
    primero = f"t.{orden} {'DESC' if desc else 'ASC'}, " if orden in propias else ""
    sql = f"SELECT {sel} FROM {d['tabla']} t {join} WHERE {w} ORDER BY {primero}{', '.join('t.' + c for c in d['llave'])} LIMIT :lim"
    with pg.connect() as c:
        filas = [dict(r) for r in c.execute(text(sql), {**p, "lim": limit}).mappings()]
        # total real (con los mismos filtros): la pantalla avisa si solo muestra una parte
        total = len(filas) if len(filas) < limit else c.execute(text(f"SELECT count(*) FROM {d['tabla']} t {join} WHERE {w}"), p).scalar()
        response.headers["X-Total-Rows"] = str(total)
        if nombre == "items" and filas:
            # una columna por plan de criterios (CATEGORIA, LINEA, ...) con el nombre del criterio como valor
            planes = c.execute(text("SELECT ic.plan_id, coalesce(pl.nombre, ic.plan_id) FROM (SELECT DISTINCT plan_id FROM ddmrp.item_criterio) ic "
                                    "LEFT JOIN ddmrp.criterio_plan pl ON pl.plan_id = ic.plan_id ORDER BY ic.plan_id")).all()
            nom = {pid: n.strip().upper() for pid, n in planes}
            vals = {}
            for pr, pid, cid, cn in c.execute(text("SELECT ic.producto_id, ic.plan_id, ic.criterio_id, cr.nombre FROM ddmrp.item_criterio ic "
                                                    "LEFT JOIN ddmrp.criterio cr ON cr.plan_id = ic.plan_id AND cr.criterio_id = ic.criterio_id")):
                vals.setdefault(pr, {})[pid] = (cn or cid).strip()
            for f in filas:
                v = vals.get(f["producto_id"], {})
                for pid, n in nom.items():
                    f[n] = v.get(pid)
        return filas


class Cambio(BaseModel):
    llave: dict[str, str]
    campo: str
    valor: float | int | str | None


@router.patch("/{nombre}")
def editar(nombre: str, cambios: list[Cambio], user=Depends(requiere("cargar_maestros", "cargar_maestros_compras"))):
    d = _def_para(nombre, user, True)
    with pg.begin() as c:
        for ch in cambios:
            if ch.campo not in d["editable"]:
                raise HTTPException(400, f"Campo no editable: {ch.campo}")
            if set(ch.llave) != set(d["llave"]):
                raise HTTPException(400, f"La llave debe ser {d['llave']}")
            _autorizar(user, d, ch.llave, c)
            if nombre == "criterios-filtros" and str(ch.valor) not in ("0", "1", "0.0", "1.0"):
                raise HTTPException(400, "En_filtros solo admite 1 (sale en los filtros) o 0 (no sale)")
            try:
                ch.valor = _valor_permitido(d, ch.campo, ch.valor)
            except ValueError as e:
                raise HTTPException(400, str(e))
            cast = CAST[d["editable"][ch.campo]]
            donde = " AND ".join(f"{k} = CAST(:k_{k} AS {CAST[TIPO_LLAVE.get(k, 'txt')]})" for k in d["llave"])
            kp = {f"k_{k}": ch.llave[k] for k in d["llave"]}
            ant = c.execute(text(f"SELECT {ch.campo}::text FROM {d['tabla']} WHERE {donde}"), kp).scalar()
            r = c.execute(text(f"UPDATE {d['tabla']} SET {ch.campo} = CAST(:v AS {cast}) WHERE {donde}"),
                          {**kp, "v": ch.valor})
            if r.rowcount == 0:
                raise HTTPException(404, "Fila no encontrada")
            c.execute(text("INSERT INTO ddmrp.config_audit (tabla, llave, campo, valor_anterior, valor_nuevo, usuario) "
                           "VALUES (:t, :l, :c, :a, :n, :u)"),
                      {"t": nombre, "l": "|".join(ch.llave[k] for k in d["llave"]), "c": ch.campo,
                       "a": ant, "n": None if ch.valor is None else str(ch.valor), "u": user["nombre"]})
    return {"ok": True, "cambios": len(cambios)}


# ---------------------------------------------------------------------------------------------------------
# Carga masiva (pegar / Excel / CSV) y descarga a Excel
# ---------------------------------------------------------------------------------------------------------
MAX_FILAS = 1_000_000


TODAS = {"*", "TODAS", "TODO", "TODOS", "ALL"}
NXP = {"nxp", "n x p", "condenado", "condenar"}
EMPRESA_ALIAS = {"funat": "05", "stores": "21", "store": "21", "tiendas": "21"}


def _valor_llave(col: str, v) -> str:
    """Normaliza una celda llave. Excel convierte '019901' en 19901: se recuperan los ceros."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    t = str(v).strip()
    if col == "empresa" and t and not t.isdigit():   # nombres en vez de codigo: Funat -> 05, Stores -> 21
        t = EMPRESA_ALIAS.get(t.lower(), t)
        if not t.isdigit():
            raise ValueError(f"empresa {t!r} no reconocida: usa 05 (Funat) o 21 (Stores)")
    if isinstance(v, int) and col in ANCHO:
        t = t.zfill(ANCHO[col])
    return t


def _valor_num(v, tipo: str):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, str):
        v = v.strip().replace(" ", "")
        if v.endswith("%"):                              # 20% -> 0.2 (las cotas se guardan como fraccion del TOG)
            return _valor_num(v[:-1], tipo) / 100 if tipo != "int" else int(float(v[:-1].replace(",", ".")))
        if v.count(",") == 1 and v.count(".") == 0:   # 1,5 -> 1.5
            v = v.replace(",", ".")
    x = float(v)
    return int(x) if tipo == "int" else x


def _leer_archivo(nombre_archivo: str, datos: bytes) -> list[dict]:
    """Primera hoja del Excel (o CSV): fila 1 = encabezados. Devuelve filas como dict {encabezado_en_minusculas: valor}."""
    if nombre_archivo.lower().endswith((".csv", ".txt")):
        txt = datos.decode("utf-8-sig", errors="replace")
        sep = max(",;\t", key=lambda c: txt.split("\n", 1)[0].count(c))
        filas = list(csv.reader(io.StringIO(txt), delimiter=sep))
    else:
        try:
            wb = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
        except Exception:
            raise HTTPException(400, "No pude leer el archivo: debe ser .xlsx o .csv")
        filas = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    filas = [r for r in filas if r and any(c not in (None, "") for c in r)]
    if len(filas) < 2:
        raise HTTPException(400, "El archivo no tiene datos (fila 1 = encabezados, desde la fila 2 los datos)")
    cab = [str(c).strip().lower() if c is not None else "" for c in filas[0]]
    return [{c: (r[i] if i < len(r) else None) for i, c in enumerate(cab) if c} for r in filas[1:]]


BOD_COD = ("bodega_id", "bodega", "codigo_bodega", "cod_bodega")
BOD_NOM = ("bodega_nombre", "nombre_bodega", "tienda", "nombre")


_SI = {"1", "si", "sí", "s", "true", "x", "activo"}
_NO = {"0", "no", "n", "false", "inactivo"}
_EMAIL = re.compile(r"^[^@\s;,]+@[^@\s;,]+\.[^@\s;,]+$")


def _valor_permitido(d: dict, campo: str, v):
    """Columnas con una lista de valores permitidos (p.ej. el redondeo del proveedor): en minusculas y sin tildes; vacio = None; otro valor = error."""
    ok = d.get("valores", {}).get(campo)
    if not ok:
        return v
    t = "" if v is None else _norm(v)
    if t == "":
        return None
    if t not in ok:
        raise ValueError(f"{v!r} no es valido para {campo}: usa " + ", ".join(ok))
    return t


def _emails(t: str) -> str:
    """Varios correos separados por ; o coma (o espacios): se validan uno a uno y quedan en minusculas, sin repetidos, separados por '; '."""
    vistos: list[str] = []
    for x in re.split(r"[;,\s]+", str(t).strip()):
        if not x:
            continue
        x = x.lower()
        if not _EMAIL.match(x):
            raise ValueError(f"correo invalido: {x!r}")
        if x not in vistos:
            vistos.append(x)
    return "; ".join(vistos)


def _norm(t) -> str:
    t = unicodedata.normalize("NFKD", str(t or "")).encode("ascii", "ignore").decode().lower()
    return " ".join(t.split())


def _mapa_bodegas():
    """(por codigo, por nombre) desde ddmrp.maestro_bodega (que el ETL trae de Siesa t150)."""
    with pg.connect() as c:
        rows = c.execute(text("SELECT bodega_id, empresa, descripcion FROM ddmrp.maestro_bodega")).all()
    por_cod: dict[str, list[tuple[str, str]]] = {}
    for r in rows:   # un mismo codigo puede existir en las dos empresas (Melonn: MBG01 en Funat y en Stores)
        por_cod.setdefault(r[0].strip().upper(), []).append((r[0], r[1]))
    return (por_cod, {_norm(r[2]): (r[0], r[1]) for r in rows if r[2]}, sorted({r[1] for r in rows}))


def _resolver_tienda(cands: list[tuple[str, str]], emp_arch: str):
    """Una tienda por su codigo: si el codigo esta en las dos empresas, la empresa del archivo decide; sin ella no se adivina."""
    if emp_arch:
        hit = [c for c in cands if c[1] == emp_arch]
        return hit[0] if hit else None
    if len(cands) > 1:
        raise ValueError(f"la tienda {cands[0][0]} existe en las empresas {' y '.join(sorted(c[1] for c in cands))}: agrega la columna empresa")
    return cands[0]


def _preparar(d: dict, filas: list[dict], omitir: bool = False, malas_out: list | None = None, cond_out: list | None = None) -> tuple[list[str], list[dict]]:
    """Valida y normaliza. La tienda se identifica por codigo (SV021) o por nombre y la empresa se deduce de la tienda.
    Devuelve (columnas editables presentes, filas listas)."""
    presentes = set().union(*[f.keys() for f in filas]) if filas else set()
    cod_h = next((h for h in BOD_COD if h in presentes), None)
    nom_h = next((h for h in BOD_NOM if h in presentes), None)
    if d.get("nombre_en_tabla") and nom_h:   # bitacora: la columna del nombre (bodega_nombre / tienda / nombre) tambien se guarda como descripcion
        presentes = presentes | {"descripcion"}
    con_bodega = bool(d.get("bodega"))
    otras = [k for k in d["llave"] if k != "empresa" and not (con_bodega and k == d["bodega"])]
    faltan = [k for k in otras if k not in presentes]
    if con_bodega and not (cod_h or nom_h):
        faltan.insert(0, "bodega_id (codigo, ej. SV021) o bodega_nombre")
    if faltan:
        amigable = ([d["bodega"]] if con_bodega else []) + otras + list(d["editable"])
        raise HTTPException(400, f"Faltan columnas: {faltan}. Encabezados: {amigable}")
    cols = [c for c in d["editable"] if c in presentes]
    # "NxP" (condenar) en la columna vectorizacion: la referencia no se despacha a esa tienda (va al maestro de condenados)
    es_nxp = lambda v: isinstance(v, str) and v.strip().lower() in NXP
    por_cod, por_nom, empresas = _mapa_bodegas() if ("empresa" in d["llave"] or con_bodega) else ({}, {}, [])
    malas, out = [], []
    cedis_ok: dict[str, set] = {}
    if d.get("valida_cedi") and "cedi" in cols:
        with pg.connect() as c:
            for emp_c, bod_c in c.execute(text("SELECT empresa, bodega_id FROM ddmrp.maestro_cedi")).all():
                cedis_ok.setdefault(bod_c.strip().upper(), set()).add(emp_c)
    for n, f in enumerate(filas, start=2):
        try:
            if d.get("nombre_en_tabla") and nom_h and not f.get("descripcion"):
                f = {**f, "descripcion": f.get(nom_h)}
            llave: dict[str, str] = {}
            if con_bodega:
                cod = _valor_llave("bodega_id", f.get(cod_h)) if cod_h else ""
                emp_f = _valor_llave("empresa", f.get("empresa")) if "empresa" in d["llave"] else ""
                hit = _resolver_tienda(por_cod[cod.upper()], emp_f) if cod and cod.upper() in por_cod else None
                if d.get("todas") and not hit and (cod.upper() in TODAS or (not cod and not (nom_h and f.get(nom_h)))):
                    hit = (d["todas"], None)   # el item queda condenado en todas las tiendas
                if not hit and (cod or nom_h):
                    hit = por_nom.get(_norm(cod)) or (por_nom.get(_norm(f.get(nom_h))) if nom_h else None)
                if not hit:
                    if cod and (not por_cod or _valor_llave("empresa", f.get("empresa"))):
                        hit = (cod, None)   # maestro vacio o empresa indicada en el archivo: se acepta el codigo tal cual
                    else:
                        raise ValueError(f"tienda no reconocida: {cod or f.get(nom_h)!r}" + (
                            ". Si es una tienda NUEVA (no esta en el catalogo de Siesa) escribe tambien la empresa: 05 Funat o 21 Stores" if cod else ""))
                llave[d["bodega"]] = hit[0]
                if "empresa" in d["llave"]:
                    emp_arch = _valor_llave("empresa", f.get("empresa"))
                    # el mismo codigo de bodega puede existir en las dos empresas (Melonn: MBG01 en Funat y en Stores): manda la empresa del archivo
                    emp = emp_arch or hit[1] or ""
                    if not emp and not (d.get("todas") and hit[0] == d["todas"]):
                        raise ValueError("no se pudo deducir la empresa de la tienda " + repr(hit[0]) + ": esta base aun no conoce las tiendas. "
                                     "Agrega la columna empresa (05 Funat, 21 Store) al archivo, o carga primero la bitacora con esa columna")
                    llave["empresa"] = emp   # vacio solo en "todas las tiendas": mas abajo se replica en cada empresa
            for k in otras:
                llave[k] = _valor_llave(k, f.get(k))
                if TIPO_LLAVE.get(k) == "int" and llave[k]:
                    llave[k] = str(int(float(llave[k])))
            if "empresa" in d["llave"] and "empresa" not in llave:   # tabla por producto: sin empresa = todas
                emp = _valor_llave("empresa", f.get("empresa"))
                llave["empresa"] = emp
            if not all(v for k, v in llave.items() if k != "empresa"):
                raise ValueError("llave vacia")
            fila_base = dict(llave)
            cond_nxp = bool(d.get("nxp")) and es_nxp(f.get("vectorizacion"))
            if cond_nxp and cond_out is not None:
                cond_out.append((llave["empresa"], llave[d["bodega"]], llave[d["producto"]]))
            for c in cols:
                t = d["editable"][c]
                if cond_nxp and c == "vectorizacion":
                    fila_base[c] = 0
                    continue
                if c in d.get("booleanas", ()):   # activo: 1 / 0 / si / no (vacio = None: no cambia)
                    v = f.get(c)
                    v = "" if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)).strip().lower()
                    if v not in _SI | _NO | {""}:
                        raise ValueError(f"{f.get(c)!r} no es valido para {c}: usa 1/0 o si/no")
                    fila_base[c] = None if v == "" else (1 if v in _SI else 0)
                    continue
                fila_base[c] = _valor_num(f.get(c), t) if t in ("num", "int") else (str(f[c]).strip() if f.get(c) not in (None, "") else None)
                if c == "cedi" and fila_base[c]:
                    fila_base[c] = fila_base[c].upper()
                if c in d.get("emails", ()) and fila_base[c]:
                    fila_base[c] = _emails(fila_base[c])
                if c in d.get("valores", {}):
                    fila_base[c] = _valor_permitido(d, c, fila_base[c])
                if d.get("fraccion") and fila_base[c] is not None and not 0 <= fila_base[c] <= 1:
                    raise ValueError(f"{f.get(c)!r} no es un porcentaje valido: escribe 20% (o 0.2)")
            if d.get("valida_cedi") and fila_base.get("cedi"):
                emp_c = fila_base.get("empresa")
                if not any(not emp_c or e == emp_c for e in cedis_ok.get(fila_base["cedi"], ())):
                    raise ValueError(f"el CEDI {fila_base['cedi']!r} no esta en el maestro CEDIs" + (f" para la empresa {emp_c}" if emp_c else "") + ": crealo primero ahi")
            if d.get("empresa_obligatoria") and not fila_base.get("empresa"):
                raise ValueError("falta la empresa (05 Funat, 21 Store)")
            if "empresa" in d["llave"] and not fila_base.get("empresa"):
                if not empresas:
                    raise ValueError("falta la columna empresa (y aun no hay bodegas cargadas para deducirla)")
                for emp in empresas:
                    out.append({**fila_base, "empresa": emp})
            else:
                out.append(fila_base)
        except (ValueError, TypeError) as e:
            malas.append(f"fila {n}: {e}")
    if malas and omitir:
        if malas_out is not None:
            malas_out.extend(malas)
    elif malas:
        raise HTTPException(400, f"{len(malas)} fila(s) invalidas, p.ej.: " + "; ".join(malas[:5]))
    if len(out) > MAX_FILAS:
        raise HTTPException(400, f"Maximo {MAX_FILAS:,} filas por carga")
    return cols, out


def _autorizar_lote(user, d: dict, filas: list[dict], conn, reemplazo: bool):
    if reemplazo and not _puede_global(user):
        raise HTTPException(403, "Reemplazar un maestro completo requiere admin o permiso de edicion sobre todas las bodegas")
    if d.get("bodega"):
        cats: dict[str, str | None] = {}
        if d.get("producto"):
            prods = sorted({f[d["producto"]] for f in filas})
            cats = dict(conn.execute(text("SELECT producto_id, categoria FROM ddmrp.maestro_item WHERE producto_id = ANY(:p)"),
                                     {"p": prods}).all())
        asigs = asignaciones(user["id"])
        pares = {(f[d["bodega"]], cats.get(f[d["producto"]]) if d.get("producto") else None) for f in filas}
        malos = [p for p in pares if nivel_efectivo(asigs, p[0], p[1]) < NIVELES["edicion"]]
        if malos:
            raise HTTPException(403, f"Sin permiso de edicion en {len(malos)} combinacion(es) bodega/laboratorio, p.ej. {malos[:3]}")
    elif d.get("global_"):
        if not _puede_global(user):
            raise HTTPException(403, "Requiere permiso de edicion sobre todas las bodegas")
    elif not _es_admin(user):
        raise HTTPException(403, "Solo administradores")


def _importar(nombre: str, d: dict, filas: list[dict], modo: str, user, omitir: bool = False) -> dict:
    try:
        res = _importar_(nombre, d, filas, modo, user, omitir)
        bitacora.evento("carga", bitacora.usuario_de(user), f"{d['label']}: {modo}, {res.get('filas')} filas "
                        f"({res.get('creadas', 0)} nuevas, {res.get('actualizadas', 0)} actualizadas, {res.get('omitidas', 0)} omitidas)",
                        {"maestro": nombre, "modo": modo, **{k: res.get(k) for k in ("filas", "unicas", "repetidas", "creadas", "actualizadas", "omitidas")}})
        return res
    except DBAPIError as e:  # p.ej. NOT NULL o valor fuera de rango: es error del archivo, no del servidor
        raise HTTPException(400, "El archivo tiene datos que la base rechaza: " + str(e.orig).splitlines()[0][:300])


def _importar_(nombre: str, d: dict, filas: list[dict], modo: str, user, omitir: bool = False) -> dict:
    if modo not in ("actualizar", "reemplazar"):
        raise HTTPException(400, "modo debe ser 'actualizar' o 'reemplazar'")
    reemplazo = modo == "reemplazar"
    if reemplazo and d.get("sin_reemplazo"):
        raise HTTPException(400, "Este maestro no admite reemplazo total: usa 'actualizar'")
    malas: list[str] = []
    cond_nxp: list[tuple[str, str, str]] = []
    cols, filas = _preparar(d, filas, omitir, malas, cond_nxp)
    if not filas:
        raise HTTPException(400, "El archivo no trae filas")
    llave, tabla = d["llave"], d["tabla"]
    defecto = d.get("defecto", {})
    with pg.begin() as c:
        _autorizar_lote(user, d, filas, c, reemplazo)
        todas = llave + cols
        c.execute(text("CREATE TEMP TABLE stg_carga (_n serial, " + ", ".join(f"{k} text" for k in todas) + ") ON COMMIT DROP"))
        raw = c.connection
        with raw.cursor() as cur, cur.copy(f"COPY stg_carga ({', '.join(todas)}) FROM STDIN") as cp:
            for f in filas:
                cp.write_row([None if f[k] is None else str(f[k]) for k in todas])
        uniq = (f"(SELECT DISTINCT ON ({', '.join(llave)}) * FROM stg_carga ORDER BY {', '.join(llave)}, _n DESC) s")
        tipos = {**{k: TIPO_LLAVE.get(k, "txt") for k in llave}, **d["editable"]}
        def expr(k):
            e = f"CAST(NULLIF(s.{k}, '') AS {CAST[tipos[k]]})"
            return f"COALESCE({e}, s.{defecto[k]})" if k in defecto and k in cols else e
        ins_cols = list(todas) + [k for k in defecto if k not in cols]
        ins_val = [expr(k) if k in todas else f"s.{defecto[k]}" for k in ins_cols]
        if d.get("solo_existentes"):
            sel_where = f"WHERE EXISTS (SELECT 1 FROM {tabla} t WHERE " + " AND ".join(f"t.{k} = s.{k}" for k in llave) + ")"
        else:
            sel_where = ""
        n_unicas = c.execute(text(f"SELECT count(*) FROM {uniq}")).scalar()
        existentes = c.execute(text(f"SELECT count(*) FROM {uniq} JOIN {tabla} t ON " +
                                    " AND ".join(f"t.{k} = CAST(s.{k} AS {CAST[tipos[k]]})" for k in llave))).scalar()
        if reemplazo:
            c.execute(text(f"DELETE FROM {tabla}"))
            existentes = 0
        conservar = d.get("conservar", ())
        upd = ", ".join((f"{k} = COALESCE(EXCLUDED.{k}, {tabla}.{k})" if k in conservar else f"{k} = EXCLUDED.{k}") for k in cols)
        c.execute(text(f"INSERT INTO {tabla} ({', '.join(ins_cols)}) SELECT {', '.join(ins_val)} FROM {uniq} {sel_where} "
                       f"ON CONFLICT ({', '.join(llave)}) " + (f"DO UPDATE SET {upd}" if upd else "DO NOTHING")))
        if d.get("solo_existentes"):
            ignoradas = n_unicas - existentes
            creadas = 0
        else:
            ignoradas, creadas = 0, n_unicas - existentes
        if cond_nxp:
            c.execute(text("INSERT INTO ddmrp.item_condenado (empresa, bodega_id, producto_id, motivo) SELECT * FROM unnest(CAST(:e AS text[]), CAST(:b AS text[]), CAST(:p AS text[]), CAST(:m AS text[])) "
                           "ON CONFLICT (empresa, producto_id, bodega_id) DO NOTHING"),
                      {"e": [x[0] for x in cond_nxp], "b": [x[1] for x in cond_nxp], "p": [x[2] for x in cond_nxp], "m": ["NxP (carga de vectorizacion)"] * len(cond_nxp)})
        c.execute(text("INSERT INTO ddmrp.config_audit (tabla, llave, campo, valor_anterior, valor_nuevo, usuario) "
                       "VALUES (:t, '*', :m, NULL, :n, :u)"),
                  {"t": nombre, "m": "carga_" + modo, "n": f"{len(filas)} filas ({creadas} nuevas)", "u": user["nombre"]})
    return {"ok": True, "modo": modo, "filas": len(filas), "unicas": n_unicas, "repetidas": len(filas) - n_unicas, "creadas": creadas,
            "actualizadas": 0 if reemplazo else existentes, "ignoradas": ignoradas, "condenados_nxp": len(cond_nxp),
            "omitidas": len(malas), "ejemplos_omitidas": malas[:5]}


class Carga(BaseModel):
    filas: list[dict[str, float | int | str | None]]
    modo: str = "actualizar"
    omitir_invalidas: bool = False


@router.post("/{nombre}/carga")
def cargar(nombre: str, body: Carga, user=Depends(requiere("cargar_maestros", "cargar_maestros_compras"))):
    """Alta/actualizacion masiva (pegar desde Excel). Cada fila trae las columnas llave y las editables."""
    d = _def_para(nombre, user, True)
    return _importar(nombre, d, [{str(k).strip().lower(): v for k, v in f.items()} for f in body.filas], body.modo, user, body.omitir_invalidas)


@router.post("/{nombre}/excel")
async def subir_excel(nombre: str, archivo: UploadFile = File(...), modo: str = Query("actualizar"),
                      omitir_invalidas: bool = Query(False), user=Depends(requiere("cargar_maestros", "cargar_maestros_compras"))):
    """Carga un .xlsx / .csv. modo=actualizar: crea las filas nuevas y actualiza las existentes (solo las columnas
    que vengan en el archivo). modo=reemplazar: borra el maestro y deja exactamente lo del archivo."""
    d = _def_para(nombre, user, True)
    datos = await archivo.read()
    if len(datos) > 60_000_000:
        raise HTTPException(400, "Archivo demasiado grande")
    return _importar(nombre, d, _leer_archivo(archivo.filename or "", datos), modo, user, omitir_invalidas)


@router.get("/{nombre}/excel")
def bajar_excel(nombre: str, bodega_id: str | None = None, q: str | None = None, lab: str | None = None,
                empresa: str | None = None, mes: str | None = None, estado: str | None = None,
                col: str | None = None, op: str | None = None, val: str | None = None,
                categoria: str | None = None, abastecimiento: str | None = None,
                user=Depends(requiere("ver_maestros", "cargar_maestros", "ver_compras", "cargar_maestros_compras"))):
    """Descarga lo que el usuario puede ver, con los mismos encabezados que acepta la carga (sirve de plantilla).
    Maestros por tienda: la tienda va con su codigo y su nombre (la empresa se deduce). Solo trae las filas que ya estan en el maestro:
    nada se agrega solo (para sumar una tienda a la bitacora se escribe o se pega)."""
    d = _def_para(nombre, user, False)
    bod = d.get("bodega")
    llave_xl = [k for k in d["llave"] if not (bod and k in ("empresa", bod))]
    con_emp = bool(bod) and "empresa" in d["llave"]      # la empresa viaja en el archivo: asi se puede subir a otra base sin tiendas cargadas
    cols_t = [k for k in d["llave"] if k != "empresa"] if bod else d["llave"]
    sel = ", ".join(f"t.{c}" for c in cols_t + list(d["editable"]))
    join, w, p = "", "TRUE", {}
    wf, pf = _filtros_maestro(nombre, d, bodega_id, q, lab, empresa, mes, estado, col, op, val)   # los mismos filtros de la pantalla: se baja lo filtrado (sin tope de 20.000)
    if d.get("producto"):
        join = f"LEFT JOIN ddmrp.maestro_item mi ON mi.producto_id = t.{d['producto']}"
    if bod:
        join = f"LEFT JOIN ddmrp.maestro_item mi ON mi.producto_id = t.{d['producto']}" if d.get("producto") else ""
        join += f" LEFT JOIN ddmrp.maestro_bodega mb ON mb.bodega_id = t.{bod}" + (" AND mb.empresa = t.empresa" if "empresa" in d["llave"] else "")
        w, p = filtro_sql_permisos(user, f"t.{bod}", "mi.categoria" if d.get("producto") else "'*'")
        nom_sel = "" if d.get("nombre_en_tabla") else ", mb.descripcion AS bodega_nombre"   # bitacora: el nombre es su columna descripcion
        sel = (f"t.{bod}{nom_sel}" + (", t.empresa" if con_emp else "") + "".join(f", t.{c}" for c in cols_t if c != bod)
               + "".join(f", t.{c}" for c in d["editable"]))
        cab = [bod] + ([] if d.get("nombre_en_tabla") else ["bodega_nombre"]) + (["empresa"] if con_emp else []) + [c for c in cols_t if c != bod] + list(d["editable"])
    else:
        cab = cols_t + list(d["editable"])
    if d.get("lectura"):    # items-compras: columnas informativas (la carga las ignora), mismos items que la pantalla
        j, s_, cab_l = _lectura_compras(); join += j; sel += ", mi.descripcion AS descripcion" + s_; cab = cab + cab_l
        wf += " AND EXISTS (SELECT 1 FROM ddmrp.item_empresa ie WHERE ie.empresa = t.empresa AND ie.producto_id = t.producto_id) AND " + solo_activos("t")
        fw, fp = _filtro_lectura(categoria, abastecimiento); wf += fw; pf.update(fp)
    if nombre == "items":   # costo promedio: columna informativa (la carga la ignora)
        join += " LEFT JOIN ddmrp.costo_item ci ON ci.producto_id = t.producto_id"
        sel += ", ci.costo::float8 AS costo, " + "(SELECT CASE ix2.estado WHEN 1 THEN \'Activo\' WHEN 0 THEN \'Inactivo\' WHEN 2 THEN \'Bloqueado\' END FROM ddmrp.item_estado ix2 WHERE ix2.empresa = :emp AND ix2.producto_id = t.producto_id) AS estado"; cab = cab + ["costo", "estado"]
    with pg.connect() as c:
        w, p = w + wf, {**p, **pf}
        filas = c.execute(text(f"SELECT {sel} FROM {d['tabla']} t {join} WHERE {w} ORDER BY "
                               + ", ".join("t." + k for k in d["llave"]) + " LIMIT 1000000"), p).all()
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Datos")
    ws.append(cab)
    for r in filas:
        ws.append([float(v) if hasattr(v, "as_tuple") else v for v in r])
    ins = wb.create_sheet("Instrucciones")
    lineas = [d["label"], "", "Fila 1 = encabezados (no los cambies). Desde la fila 2, los datos."]
    if bod:
        lineas += ["Tienda: usa el codigo de Siesa (bodega_id, ej. SV021) o, si prefieres, solo el nombre (bodega_nombre). La empresa se deduce sola.",
                   ("La columna descripcion es el NOMBRE de la tienda y se guarda en la bitacora (tambien se acepta con el encabezado bodega_nombre, tienda o nombre). Celda vacia = no cambia el nombre."
                    if d.get("nombre_en_tabla") else "La columna bodega_nombre es informativa."),
                   "Empresa (05 Funat, 21 Store): si la columna viene llena se usa tal cual; si no, se deduce de la tienda (requiere tener las tiendas cargadas)."]
    elif "empresa" in d["llave"]:
        lineas.append("Empresa (05 Funat, 21 Store): opcional; si la dejas vacia aplica a todas.")
    otras = [k for k in d["llave"] if k not in ("empresa",) and k != bod]
    if otras:
        lineas.append("Otras columnas llave: " + ", ".join(otras))
    lineas += ["Columnas a cargar: " + (", ".join(d["editable"]) or "(ninguna: la existencia de la fila es el dato)"),
               "Actualizar: crea las filas nuevas y actualiza las existentes; solo toca las columnas que vengan en el archivo.",
               "Reemplazar: borra el maestro y deja exactamente lo del archivo (solo admin / edicion en todas las bodegas)."]
    if "producto_id" in d["llave"]:
        lineas.append("Las referencias van como texto de 6 digitos (ej. 019901).")
    for l in lineas:
        ins.append([l])
    buf = io.BytesIO(); wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{nombre}.xlsx"'})


MAX_ELIMINAR = 200_000


class Eliminar(BaseModel):
    llaves: list[dict[str, str]]


@router.post("/{nombre}/eliminar")
def eliminar(nombre: str, body: Eliminar, user=Depends(requiere("cargar_maestros", "cargar_maestros_compras"))):
    """Borra filas por su llave (p.ej. quitar una tienda de un DIA), de a miles en una sola sentencia. No aplica a factores ni parametros.
    Para borrar todo lo que cumple un filtro (sin pasar por la seleccion de la pantalla) esta /vaciar."""
    d = _def_para(nombre, user, True)
    if d.get("sin_reemplazo") or d.get("sin_eliminar") or nombre == "factores":
        raise HTTPException(400, "Este maestro no admite eliminar filas")
    if len(body.llaves) > MAX_ELIMINAR:
        raise HTTPException(400, f"Maximo {MAX_ELIMINAR:,} filas por eliminacion: para mas usa 'Eliminar todo lo filtrado'")
    if not body.llaves:
        return {"ok": True, "eliminadas": 0}
    for ll in body.llaves:
        if set(ll) != set(d["llave"]):
            raise HTTPException(400, f"La llave debe ser {d['llave']}")
    cols = d["llave"]
    with pg.begin() as c:
        _autorizar_lote(user, d, body.llaves, c, False)   # permisos de todo el lote con una sola consulta
        arr = ", ".join(f"CAST(:a{i} AS text[])" for i in range(len(cols)))
        als = ", ".join(f"c{i}" for i in range(len(cols)))
        donde = " AND ".join(f"t.{k} = CAST(s.c{i} AS {CAST[TIPO_LLAVE.get(k, 'txt')]})" for i, k in enumerate(cols))
        n = c.execute(text(f"DELETE FROM {d['tabla']} t USING unnest({arr}) AS s({als}) WHERE {donde}"),
                      {f"a{i}": [ll[k] for ll in body.llaves] for i, k in enumerate(cols)}).rowcount
        c.execute(text("INSERT INTO ddmrp.config_audit (tabla, llave, campo, valor_anterior, valor_nuevo, usuario) "
                       "VALUES (:t, '*', 'eliminar', :n, NULL, :u)"), {"t": nombre, "n": f"{n} filas", "u": user["nombre"]})
    return {"ok": True, "eliminadas": n}


class Vaciar(BaseModel):
    bodega_id: str | None = None
    q: str | None = None
    lab: str | None = None
    empresa: str | None = None
    mes: str | None = None
    col: str | None = None
    op: str | None = None
    val: str | None = None
    solo_contar: bool = False   # True: solo dice cuantas filas se borrarian (para pedir confirmacion)
    todo: bool = False          # True: confirma que se vacia el maestro completo (sin ningun filtro)


@router.post("/{nombre}/vaciar")
def vaciar(nombre: str, body: Vaciar, user=Depends(requiere("cargar_maestros", "cargar_maestros_compras"))):
    """Borra TODAS las filas que cumplen los filtros de la pantalla (no solo las que se ven ni las seleccionadas), sin tope de filas.
    Sin filtros vacia el maestro completo y exige todo=true. Es la forma de dejar un maestro limpio para volverlo a subir."""
    d = _def_para(nombre, user, True)
    if d.get("sin_reemplazo") or nombre in SIN_VACIAR:
        raise HTTPException(400, "Este maestro no se vacia desde aqui")
    if not _puede_global(user):
        raise HTTPException(403, "Vaciar por filtro requiere admin o permiso de edicion sobre todas las bodegas y categorias; con permisos parciales usa 'Eliminar seleccionadas'")
    wf, pf = _filtros_maestro(nombre, d, body.bodega_id, body.q, body.lab, body.empresa, body.mes, None, body.col, body.op, body.val)
    if not wf and not body.solo_contar and not body.todo:
        raise HTTPException(400, "Sin filtros se vacia el maestro completo: confirma con todo=true")
    prod = d.get("producto")
    join = f"LEFT JOIN ddmrp.maestro_item mi ON mi.producto_id = t.{prod}" if prod else ""
    base = f"FROM {d['tabla']} t {join} WHERE TRUE{wf}"
    with pg.begin() as c:
        n = c.execute(text(f"SELECT count(*) {base}"), pf).scalar()
        if body.solo_contar:
            return {"ok": True, "filas": n, "con_filtro": bool(wf)}
        n = c.execute(text(f"DELETE FROM {d['tabla']} WHERE ctid IN (SELECT t.ctid {base})"), pf).rowcount
        desc = ", ".join(f"{k}={v}" for k, v in body.model_dump(exclude={"solo_contar", "todo"}).items() if v) or "sin filtros (todo)"
        c.execute(text("INSERT INTO ddmrp.config_audit (tabla, llave, campo, valor_anterior, valor_nuevo, usuario) "
                       "VALUES (:t, '*', 'vaciar', :n, :f, :u)"), {"t": nombre, "n": f"{n} filas", "f": desc[:500], "u": user["nombre"]})
    bitacora.evento("carga", bitacora.usuario_de(user), f"{d['label']}: vaciado ({desc[:200]}), {n} filas eliminadas", {"maestro": nombre, "eliminadas": n})
    return {"ok": True, "eliminadas": n, "con_filtro": bool(wf)}
