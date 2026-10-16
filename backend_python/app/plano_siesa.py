"""Archivo plano de Siesa: requisicion inter-bodega (RIT) desde el CEDI a cada tienda.
Formato tomado de docs/Abastecimiento dia 4_15102026.xlsx: hojas Inicial, Documentos, Movimientos y Final.
Un documento por bodega destino; un movimiento por referencia. Las referencias van como texto (ceros a la izquierda)."""
from datetime import date
from io import BytesIO
from itertools import groupby
from openpyxl import Workbook
from .config import settings

MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE",
         "NOVIEMBRE", "DICIEMBRE"]

DOC_COLS = ["COMPAÑIA", "f440_id_co", "f440_id_tipo_docto", "f440_consec_docto", "f440_fecha", "f440_id_solicitante",
            "f440_fecha_entrega", "f440_num_dias_entrega", "f440_notas", "f440_id_bodega_salida",
            "f440_id_bodega_entrada", "f440_referencia"]
MOV_COLS = ["COMPAÑIA", "f441_id_co", "f441_id_tipo_docto", "f441_consec_docto", "f441_nro_registro",
            "f441_referencia_item", "f441_id_bodega", "f441_id_motivo", "f441_id_unidad_medida", "f441_cant_base",
            "f441_fecha_entrega", "f441_num_dias_entrega", "f441_id_co_movto", "f441_id_un_movto"]


def notas_por_defecto(fecha: date, dia: int | None = None) -> str:
    return f"ABASTECIMIENTO DEL {fecha.day} DE {MESES[fecha.month - 1]}" + (f" DÍA {dia}" if dia else "")


def co_derivado(bodega_id: str) -> str:
    """Regla observada en el archivo de ejemplo (37 de 37): los 3 digitos finales del codigo; si no hay, '000'."""
    t = bodega_id[-3:]
    return t if t.isdigit() else "000"


def generar_xlsx(lineas: list[tuple[str, str, str, float]], *, fecha: date, consec_inicial: int = 1,
                 notas: str | None = None) -> bytes:
    """lineas = [(bodega_destino, id_co, referencia, cantidad)], cantidades > 0."""
    ymd = fecha.strftime("%Y%m%d")
    notas = notas or notas_por_defecto(fecha)
    cia = settings.siesa_cia
    wb = Workbook()
    ini = wb.active; ini.title = "Inicial"
    ini.append(["COMPAÑIA"]); ini.append(["1" if cia == 1 else str(cia)])
    docs = wb.create_sheet("Documentos"); docs.append(DOC_COLS)
    movs = wb.create_sheet("Movimientos"); movs.append(MOV_COLS)
    fin = wb.create_sheet("Final")
    fin.append(["COMPAÑIA"]); fin.append(["1" if cia == 1 else str(cia)])

    consec = consec_inicial
    for bodega, grupo in groupby(sorted(lineas, key=lambda x: (x[0], x[2])), key=lambda x: x[0]):
        grupo = list(grupo)
        co = grupo[0][1]
        docs.append([cia, co, settings.plano_tipo_docto, consec, ymd, settings.plano_solicitante, ymd,
                     settings.plano_dias_entrega, notas, settings.plano_bodega_salida, bodega, settings.plano_referencia])
        for n, (_, _, ref, cant) in enumerate(grupo, start=1):
            movs.append([cia, co, settings.plano_tipo_docto, consec, n, ref, settings.plano_bodega_salida,
                         settings.plano_motivo, settings.plano_unidad_medida, int(round(cant)), ymd,
                         settings.plano_dias_entrega, "000", settings.plano_un_movto])
        consec += 1
    for row in movs.iter_rows(min_row=2, min_col=2, max_col=2):  # columna del CO como texto, igual al ejemplo
        row[0].number_format = "@"
    out = BytesIO(); wb.save(out)
    return out.getvalue()
