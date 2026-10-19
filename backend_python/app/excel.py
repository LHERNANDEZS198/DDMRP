"""Excel de salida con openpyxl en modo write_only: aguanta cientos de miles de filas sin llenar la memoria.
Cada hoja se describe con una lista de columnas {t: titulo, k: clave o posicion, f: formato, w: ancho, zona: pinta la celda con el color de su zona}."""
import io
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_FILAS = 1_000_000          # Excel aguanta 1.048.576 filas por hoja (una es el encabezado)
# colores del semaforo: los del Excel de trabajo de Lucas (rojo rosado, amarillo naranja, verde, exceso azul)
ZONA_COLOR = {"ROJA": "E8B9B7", "AMARILLA": "FFC000", "VERDE": "92D050", "EXCESO": "4792DC"}
ZONA_LETRA = {"R": "ROJA", "A": "AMARILLA", "V": "VERDE", "E": "EXCESO"}
_FILL = {k: PatternFill("solid", start_color=c, end_color=c) for k, c in ZONA_COLOR.items()}
_ENCABEZADO = PatternFill("solid", start_color="DCE6C8", end_color="DCE6C8")


def valor(v, f: str | None = None):
    """Valor listo para la celda: Decimal -> float, formatos redondeados en Python (sin estilo por celda: mucho mas rapido)."""
    if v is None:
        return None
    if isinstance(v, Decimal):
        v = float(v)
    if isinstance(v, bool):
        return "SI" if v else ""
    if isinstance(v, (int, float)):
        if f in ("int", "cop"):
            return int(round(v))
        if f in ("dec1", "pct1"):
            return round(v, 1)
        if f == "dec2":
            return round(v, 2)
    if hasattr(v, "isoformat"):
        return v.isoformat()[:19].replace("T", " ") if hasattr(v, "hour") else v.isoformat()
    return v


def libro() -> Workbook:
    wb = Workbook(write_only=True)
    return wb


def _encabezado(ws, cols):
    cel = []
    for c in cols:
        x = WriteOnlyCell(ws, value=c["t"])
        x.font = Font(bold=True); x.fill = _ENCABEZADO
        x.alignment = Alignment(wrap_text=True, vertical="center")
        cel.append(x)
    ws.append(cel)


def hoja(wb: Workbook, titulo: str, cols: list[dict], filas, fijar: int = 0) -> int:
    """Escribe una hoja. `filas` = iterable de dict (clave c['k']) o de secuencias (posicion c['k']). Devuelve las filas de datos escritas.
    `fijar` = columnas fijas a la izquierda (el encabezado siempre queda fijo)."""
    ws = wb.create_sheet(titulo[:31])
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = c.get("w", 14)
    ws.freeze_panes = f"{get_column_letter(fijar + 1)}2"
    _encabezado(ws, cols)
    n = 0
    for r in filas:
        if n >= MAX_FILAS:
            break
        cel = []
        for c in cols:
            v = valor(r.get(c["k"]) if hasattr(r, "get") else r[c["k"]], c.get("f"))
            if c.get("zona") and v in _FILL:
                x = WriteOnlyCell(ws, value=v); x.fill = _FILL[v]; cel.append(x)
            else:
                cel.append(v)
        ws.append(cel)
        n += 1
    if n:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n + 1}"
    return n


def hoja_mapa(wb: Workbook, titulo: str, cols: list[dict], fechas: list, filas) -> int:
    """Mapa de calor: columnas fijas + una columna por fecha; cada celda lleva el numero y el color de su zona (fila['v'], fila['z'])."""
    ws = wb.create_sheet(titulo[:31])
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = c.get("w", 14)
    for j in range(len(fechas)):
        ws.column_dimensions[get_column_letter(len(cols) + j + 1)].width = 7.5
    ws.freeze_panes = f"{get_column_letter(min(len(cols), 4) + 1)}2"
    _encabezado(ws, cols + [{"t": str(f)[5:] if len(str(f)) == 10 else str(f)} for f in fechas])
    n = 0
    centro = Alignment(horizontal="center")
    for r in filas:
        if n >= MAX_FILAS:
            break
        cel = [valor(r[c["k"]], c.get("f")) for c in cols]
        for v, z in zip(r["v"], r["z"]):
            if v is None:
                cel.append(None); continue
            x = WriteOnlyCell(ws, value=int(round(v)))
            if z in ZONA_LETRA:
                x.fill = _FILL[ZONA_LETRA[z]]
            x.alignment = centro
            cel.append(x)
        ws.append(cel)
        n += 1
    if n:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols) + len(fechas))}{n + 1}"
    return n


def hoja_texto(wb: Workbook, titulo: str, lineas: list, ancho: tuple = (30, 90)):
    """Hoja informativa de dos columnas (etiqueta, valor) o de una sola (texto suelto)."""
    ws = wb.create_sheet(titulo[:31])
    ws.column_dimensions["A"].width = ancho[0]; ws.column_dimensions["B"].width = ancho[1]
    for l in lineas:
        if isinstance(l, (tuple, list)):
            a = WriteOnlyCell(ws, value=l[0]); a.font = Font(bold=True)
            ws.append([a, l[1]])
        else:
            ws.append([l])


def leyenda_zonas(ws_titulo: str = "Leyenda"):
    return [("Rojo", "ROJA"), ("Amarillo", "AMARILLA"), ("Verde", "VERDE"), ("Exceso", "EXCESO")]


def a_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
